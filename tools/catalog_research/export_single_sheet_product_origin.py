from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


REPO_ROOT = Path(__file__).resolve().parents[2]
SOURCE_WORKBOOK = REPO_ROOT / "output" / "spreadsheet" / "recommendation_product_data_review.xlsx"
OUTPUT_WORKBOOK = REPO_ROOT / "output" / "spreadsheet" / "catalog_product_origin_single_sheet.xlsx"

COLLECTED_AT = "2026-04-24 21:25:00 ICT"


PROVINCE_NORMALIZATION = {
    "Đà Lạt, Việt Nam": ("Việt Nam", "Lâm Đồng", "Đà Lạt"),
    "Lâm Đồng, Việt Nam": ("Việt Nam", "Lâm Đồng", ""),
    "Long An, Việt Nam": ("Việt Nam", "Long An", ""),
    "Tiền Giang, Việt Nam": ("Việt Nam", "Tiền Giang", ""),
    "Đồng Nai, Việt Nam": ("Việt Nam", "Đồng Nai", ""),
    "Bình Dương, Việt Nam": ("Việt Nam", "Bình Dương", ""),
    "Bình Thuận, Việt Nam": ("Việt Nam", "Bình Thuận", ""),
    "Ninh Thuận, Việt Nam": ("Việt Nam", "Ninh Thuận", ""),
    "Cần Thơ, Việt Nam": ("Việt Nam", "Cần Thơ", ""),
    "Bến Tre, Việt Nam": ("Việt Nam", "Bến Tre", ""),
    "TP.HCM, Việt Nam": ("Việt Nam", "TP.HCM", ""),
    "Nghệ An, Việt Nam": ("Việt Nam", "Nghệ An", ""),
    "Tây Nguyên, Việt Nam": ("Việt Nam", "Tây Nguyên", "Vùng, chưa rõ tỉnh"),
    "Đồng Tháp, Việt Nam": ("Việt Nam", "Đồng Tháp", ""),
    "Đắk Lắk, Việt Nam": ("Việt Nam", "Đắk Lắk", ""),
    "Bình Phước, Việt Nam": ("Việt Nam", "Bình Phước", ""),
    "Nhật Bản": ("Nhật Bản", "", ""),
    "Úc": ("Úc", "", ""),
    "Mỹ": ("Mỹ", "", ""),
    "Thái Lan": ("Thái Lan", "", ""),
    "Nam Mỹ": ("Nam Mỹ", "", "Chưa rõ quốc gia cụ thể"),
    "Nam Phi": ("Nam Phi", "", ""),
    "New Zealand": ("New Zealand", "", ""),
    "Trung Quốc": ("Trung Quốc", "", ""),
    "Nhật Bản (trồng VN)": ("Việt Nam", "Bắc Giang; Tây Ninh", "Giống/nhãn Nhật, trồng tại Việt Nam theo nguồn rà soát"),
    "Hàn Quốc (trồng VN)": ("Việt Nam", "Hòa Bình", "Giống/nhãn Hàn Quốc chưa xác nhận, có nguồn trồng tại Việt Nam"),
}


MANUAL_OVERRIDES = {
    46: {
        "country": "Việt Nam",
        "province": "Bắc Giang; Tây Ninh",
        "area": "Giống/nhãn Nhật, trồng tại Việt Nam",
        "note": "Dùng tỉnh/vùng đã rà: Bắc Giang, Tây Ninh; chưa chốt vùng cung ứng thật của SKU.",
    },
    52: {
        "country": "Trung Quốc",
        "province": "Hồ Bắc/Suizhou (tham khảo); Trung Quốc",
        "area": "Nấm hương khô nhập khẩu",
        "note": "Tỉnh chỉ là vùng tham khảo từ nguồn ngành hàng; cần chứng từ nhà cung cấp nếu muốn chốt chính thức.",
    },
    53: {
        "country": "Việt Nam",
        "province": "Hòa Bình",
        "area": "Nấm linh chi đỏ trồng tại Việt Nam",
        "note": "Nhãn Hàn Quốc chưa xác nhận bằng chứng từ; tạm dùng vùng trồng Việt Nam đã rà.",
    },
    61: {"country": "Nhật Bản", "province": "Aomori", "area": "", "note": ""},
    62: {"country": "Úc", "province": "Murray Valley; Riverina; Riverland", "area": "", "note": ""},
    63: {"country": "Mỹ", "province": "California", "area": "Coachella Valley; San Joaquin Valley", "note": ""},
    64: {"country": "Việt Nam", "province": "Tiền Giang", "area": "Cái Bè - Hòa Lộc", "note": ""},
    65: {"country": "Việt Nam", "province": "Long An", "area": "Vĩnh Hưng/Tân Trụ/Long Trì", "note": "Nguồn theo dưa hấu Long An nói chung."},
    66: {"country": "Việt Nam", "province": "Đắk Lắk", "area": "", "note": "Nguồn theo bơ Đắk Lắk nói chung."},
    67: {"country": "Việt Nam", "province": "Lâm Đồng", "area": "Đà Lạt/Lạc Dương", "note": ""},
    68: {
        "country": "Việt Nam",
        "province": "Ninh Thuận",
        "area": "Ninh Sơn",
        "note": "Chuối già Nam Mỹ là giống/cultivar; dùng vùng trồng tham khảo tại Ninh Thuận từ nguồn báo chí.",
    },
    69: {"country": "Việt Nam", "province": "Bình Phước", "area": "", "note": ""},
    70: {"country": "Việt Nam", "province": "Bình Thuận", "area": "Vùng chỉ dẫn địa lý thanh long", "note": ""},
    71: {"country": "Thái Lan", "province": "Chanthaburi; Rayong; Trat", "area": "Miền Đông Thái Lan", "note": "Vùng trồng tham khảo phổ biến của măng cụt Thái Lan."},
    72: {"country": "Việt Nam", "province": "Tiền Giang", "area": "Cái Bè/Cai Lậy/Châu Thành/Tân Phước/TX Cai Lậy", "note": ""},
    73: {"country": "Mỹ", "province": "California", "area": "San Joaquin Valley", "note": ""},
    74: {"country": "Nhật Bản", "province": "Ibaraki; Shizuoka", "area": "", "note": "Tỉnh tùy lô/nhà kính, cần chốt với nhà cung cấp nếu nhập DB chính thức."},
    75: {"country": "Mỹ", "province": "California", "area": "", "note": ""},
    76: {"country": "Nam Phi", "province": "Western Cape", "area": "Ceres/Groenland/Wolseley-Tulbagh", "note": ""},
    77: {"country": "New Zealand", "province": "Bay of Plenty", "area": "Te Puke/Katikati", "note": "Vùng trồng tham khảo phổ biến của kiwi New Zealand."},
    78: {"country": "Thái Lan", "province": "Prachuap Khiri Khan", "area": "Pran Buri/Thap Sakae context", "note": "Vùng trồng dứa Thái Lan tham khảo, phù hợp MD2 theo nguồn thương mại."},
    115: {"country": "Trung Quốc", "province": "Hebei; Shandong; Jiangsu", "area": "Vùng rau/leek tham khảo", "note": "Tỉnh tham khảo từ nguồn web/ChinaDaily/nhà cung cấp; cần chứng từ lô hàng nếu chốt chính thức."},
    116: {"country": "Việt Nam", "province": "Tiền Giang", "area": "", "note": ""},
    117: {"country": "Việt Nam", "province": "Long An", "area": "", "note": ""},
    118: {"country": "Việt Nam", "province": "Long An", "area": "Vĩnh Hưng/Tân Trụ/Long Trì", "note": "Nguồn theo dưa hấu Long An nói chung."},
    119: {"country": "Việt Nam", "province": "Bình Thuận", "area": "Vùng chỉ dẫn địa lý thanh long", "note": ""},
    120: {"country": "Việt Nam", "province": "Đắk Lắk", "area": "Krông Pắc/Cư M'gar", "note": "Nguồn theo sầu riêng Đắk Lắk nói chung."},
    121: {
        "category": "Củ & rễ",
        "country": "Việt Nam",
        "province": "Vĩnh Long",
        "area": "Mang Thít",
        "note": "Catalog thiếu ProductInfo; loại đã chỉnh từ Rau lá sang Củ & rễ.",
    },
    122: {
        "country": "Việt Nam",
        "province": "Long An",
        "area": "Suy từ dòng Rau muống ProductID 1 trong catalog",
        "note": "Catalog thiếu ProductInfo; tạm dùng Long An theo sản phẩm Rau muống đã có trong catalog.",
    },
}


CATEGORY_HINTS = {
    "Rau lá": "Rau lá",
    "Rau ăn hoa / thân / mầm": "Rau ăn hoa / thân / mầm",
    "Rau ăn quả": "Rau ăn quả",
    "Củ & rễ": "Củ & rễ",
    "Nấm": "Nấm",
    "Rau thơm & gia vị": "Rau thơm & gia vị",
    "Trái cây": "Trái cây",
}


HEADERS = [
    "STT",
    "ProductID",
    "Tên sản phẩm",
    "Loại sản phẩm",
    "Quốc gia",
    "Tỉnh/Thành/Vùng",
    "Khu vực chi tiết",
    "Nội địa/Nhập khẩu",
    "Nguồn lấy tỉnh/thành",
    "Ghi chú cần kiểm tra",
    "Thời gian cập nhật",
]


def source_label(status: str | None, current_origin: str | None) -> str:
    if status and status.startswith("Partially verified"):
        return "Đã rà nguồn ngoài + catalog"
    if current_origin:
        return "Catalog hiện có"
    return "Thiếu trong catalog, bổ sung tạm từ rà soát"


def domestic_imported(country: str, current_origin: str | None, verified: str | None) -> str:
    if verified:
        value = verified.lower()
        if "import" in value:
            return "Nhập khẩu"
        if "domestic" in value or "vietnam" in value:
            return "Nội địa"
    if country == "Việt Nam":
        return "Nội địa"
    if current_origin and "trồng VN" in current_origin:
        return "Nội địa"
    return "Nhập khẩu"


def normalize_origin(product_id: int, current_origin: str | None, verified_region: str | None):
    override = MANUAL_OVERRIDES.get(product_id, {})
    if override:
        return (
            override.get("country", ""),
            override.get("province", ""),
            override.get("area", ""),
            override.get("note", ""),
        )

    if current_origin in PROVINCE_NORMALIZATION:
        country, province, area = PROVINCE_NORMALIZATION[current_origin]
        return country, province, area, ""

    if verified_region:
        return "Việt Nam", verified_region, "", "Lấy từ cột verified hiện có."

    if current_origin and ", Việt Nam" in current_origin:
        province = current_origin.replace(", Việt Nam", "").strip()
        return "Việt Nam", province, "", ""

    return current_origin or "", "", "", "Chưa có tỉnh/thành rõ trong dữ liệu gốc."


def make_rows():
    source_wb = load_workbook(SOURCE_WORKBOOK, read_only=True, data_only=True)
    source_ws = source_wb["Products Review"]
    headers = [cell.value for cell in source_ws[1]]
    header_index = {header: idx for idx, header in enumerate(headers)}
    rows = []

    for seq, row in enumerate(source_ws.iter_rows(min_row=2, values_only=True), start=1):
        product_id = int(row[header_index["ProductID"]])
        product_name = row[header_index["ProductName"]]
        category = row[header_index["CategoryName"]]
        current_origin = row[header_index["Current Origin"]]
        status = row[header_index["Review Status"]]
        verified_region = row[header_index["Verified Origin Province/Region"]]
        verified_di = row[header_index["Verified Domestic/Imported"]]
        uncertainty = row[header_index["Tôi chưa chắc chắn"]]

        override = MANUAL_OVERRIDES.get(product_id, {})
        category = override.get("category", CATEGORY_HINTS.get(category, category))
        country, province, area, note = normalize_origin(product_id, current_origin, verified_region)
        note_parts = [part for part in [note, uncertainty] if part]

        rows.append(
            [
                seq,
                product_id,
                product_name,
                category,
                country,
                province or "Chưa rõ",
                area,
                domestic_imported(country, current_origin, verified_di),
                source_label(status, current_origin),
                " | ".join(dict.fromkeys(note_parts)),
                COLLECTED_AT,
            ]
        )

    return rows


def autosize_columns(ws):
    max_widths = {
        "A": 7,
        "B": 12,
        "C": 28,
        "D": 24,
        "E": 18,
        "F": 34,
        "G": 44,
        "H": 18,
        "I": 26,
        "J": 60,
        "K": 22,
    }
    for col_idx, _ in enumerate(HEADERS, start=1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = max_widths.get(letter, 18)


def build_workbook(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Danh sách sản phẩm"

    ws.append(HEADERS)
    for row in rows:
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(bold=True, color="FFFFFF")
    thin_border = Border(bottom=Side(style="thin", color="D1D5DB"))
    domestic_fill = PatternFill("solid", fgColor="ECFDF5")
    imported_fill = PatternFill("solid", fgColor="EFF6FF")
    needs_check_fill = PatternFill("solid", fgColor="FEF3C7")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        is_imported = row[7].value == "Nhập khẩu"
        row_fill = imported_fill if is_imported else domestic_fill
        if row[5].value == "Chưa rõ" or row[5].value == "Chưa rõ tỉnh":
            row_fill = needs_check_fill
        for cell in row:
            cell.fill = row_fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin_border
        row[0].alignment = Alignment(horizontal="center", vertical="top")
        row[1].alignment = Alignment(horizontal="center", vertical="top")

    autosize_columns(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"

    table = Table(displayName="CatalogProductOriginTable", ref=f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}")
    style = TableStyleInfo(
        name="TableStyleMedium4",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    table.tableStyleInfo = style
    ws.add_table(table)

    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 34
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 42

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    return wb


def main() -> None:
    rows = make_rows()
    if len(rows) != 100:
        raise RuntimeError(f"Expected 100 products, got {len(rows)}")
    wb = build_workbook(rows)
    OUTPUT_WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_WORKBOOK)
    print(f"Saved {OUTPUT_WORKBOOK}")
    print(f"Rows: {len(rows)}")
    print(f"Generated at: {datetime.now().isoformat(timespec='seconds')}")


if __name__ == "__main__":
    main()

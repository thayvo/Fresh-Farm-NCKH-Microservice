from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
INPUT_WORKBOOK = REPO_ROOT / "output" / "spreadsheet" / "catalog_product_seasonality_single_sheet.xlsx"
OUTPUT_DELTA = REPO_ROOT / "docs" / "FreshFarmCatalogDb" / "FreshFarmCatalogDB.2026-04-25.product-seasonality.delta.sql"
OUTPUT_VERIFY = REPO_ROOT / "docs" / "FreshFarmCatalogDb" / "FreshFarmCatalogDB.2026-04-25.product-seasonality.verify.sql"


def sql_string(value: object, max_len: int | None = None) -> str:
    if value is None:
        return "NULL"
    text = str(value)
    if max_len is not None:
        text = text[:max_len]
    return "N'" + text.replace("'", "''") + "'"


def sql_int(value: object) -> str:
    if value is None or value == "":
        return "NULL"
    return str(int(value))


def sql_bit(value: object) -> str:
    return "1" if str(value).strip().lower() in {"yes", "true", "1"} else "0"


def read_rows() -> list[dict[str, object]]:
    workbook = load_workbook(INPUT_WORKBOOK, read_only=True, data_only=True)
    worksheet = workbook["Mùa vụ sản phẩm"]
    headers = [cell.value for cell in worksheet[1]]
    index = {header: column for column, header in enumerate(headers)}
    rows = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        rows.append({header: row[column] for header, column in index.items()})
    return rows


def build_values(rows: list[dict[str, object]]) -> str:
    value_lines = []
    for row in rows:
        value_lines.append(
            "        ("
            + ", ".join(
                [
                    sql_int(row["ProductID"]),
                    sql_string(row["Quốc gia"], 80),
                    sql_string(row["Tỉnh/Thành/Vùng"], 200),
                    sql_string(row["Khu vực chi tiết"], 200),
                    sql_string(row["SeasonType"], 60),
                    sql_string(row["SeasonLabel"], 180),
                    sql_int(row["StartMonth"]),
                    sql_int(row["EndMonth"]),
                    sql_string(row["PeakMonths"], 80),
                    sql_bit(row["IsYearRound"]),
                    sql_bit(row["HasPeakSeason"]),
                    sql_bit(row["IsControlledCultivation"]),
                    sql_bit(row["IsImportedSeason"]),
                    sql_bit(row["IsOffSeason"]),
                    sql_bit(row["IsPostHarvestAvailability"]),
                    sql_int(row["SeasonScoreWeight"]),
                    sql_string(row["ConfidenceLevel"], 40),
                    sql_string(row["Nguồn/Ghi chú"], 600),
                ]
            )
            + ")"
        )
    return ",\n".join(value_lines)


def build_expected_count_values(rows: list[dict[str, object]]) -> str:
    counts: dict[int, int] = {}
    for row in rows:
        product_id = int(row["ProductID"])
        counts[product_id] = counts.get(product_id, 0) + 1

    return ",\n".join(
        f"        ({product_id}, {season_rows})"
        for product_id, season_rows in sorted(counts.items())
    )


def build_delta(rows: list[dict[str, object]]) -> str:
    values_sql = build_values(rows)
    expected_counts_sql = build_expected_count_values(rows)
    return f"""SET NOCOUNT ON;
GO

IF OBJECT_ID(N'dbo.Products', N'U') IS NULL
BEGIN
    RAISERROR(N'Products chua ton tai. Hay apply schema Catalog truoc.', 16, 1);
    RETURN;
END;
GO

IF OBJECT_ID(N'dbo.ProductSeasonality', N'U') IS NULL
BEGIN
    CREATE TABLE dbo.ProductSeasonality
    (
        ProductSeasonalityID int IDENTITY(1,1) NOT NULL,
        ProductID int NOT NULL,
        Country nvarchar(80) NULL,
        ProvinceRegion nvarchar(200) NULL,
        AreaDetail nvarchar(200) NULL,
        SeasonType nvarchar(60) NOT NULL,
        SeasonLabel nvarchar(180) NOT NULL,
        StartMonth tinyint NOT NULL,
        EndMonth tinyint NOT NULL,
        PeakMonths nvarchar(80) NULL,
        IsYearRound bit NOT NULL CONSTRAINT DF_ProductSeasonality_IsYearRound DEFAULT (0),
        HasPeakSeason bit NOT NULL CONSTRAINT DF_ProductSeasonality_HasPeakSeason DEFAULT (0),
        IsControlledCultivation bit NOT NULL CONSTRAINT DF_ProductSeasonality_IsControlledCultivation DEFAULT (0),
        IsImportedSeason bit NOT NULL CONSTRAINT DF_ProductSeasonality_IsImportedSeason DEFAULT (0),
        IsOffSeason bit NOT NULL CONSTRAINT DF_ProductSeasonality_IsOffSeason DEFAULT (0),
        IsPostHarvestAvailability bit NOT NULL CONSTRAINT DF_ProductSeasonality_IsPostHarvestAvailability DEFAULT (0),
        SeasonScoreWeight int NOT NULL CONSTRAINT DF_ProductSeasonality_SeasonScoreWeight DEFAULT (50),
        ConfidenceLevel nvarchar(40) NULL,
        SourceNote nvarchar(600) NULL,
        CreatedAt datetime2(0) NOT NULL CONSTRAINT DF_ProductSeasonality_CreatedAt DEFAULT (SYSUTCDATETIME()),
        UpdatedAt datetime2(0) NOT NULL CONSTRAINT DF_ProductSeasonality_UpdatedAt DEFAULT (SYSUTCDATETIME()),
        CONSTRAINT PK_ProductSeasonality PRIMARY KEY CLUSTERED (ProductSeasonalityID),
        CONSTRAINT FK_ProductSeasonality_Products_ProductID FOREIGN KEY (ProductID)
            REFERENCES dbo.Products(ProductID),
        CONSTRAINT CK_ProductSeasonality_StartMonth CHECK (StartMonth BETWEEN 1 AND 12),
        CONSTRAINT CK_ProductSeasonality_EndMonth CHECK (EndMonth BETWEEN 1 AND 12),
        CONSTRAINT CK_ProductSeasonality_SeasonScoreWeight CHECK (SeasonScoreWeight BETWEEN 0 AND 100)
    );

    CREATE UNIQUE INDEX UX_ProductSeasonality_Product_Season
        ON dbo.ProductSeasonality(ProductID, SeasonLabel, StartMonth, EndMonth);

    CREATE INDEX IX_ProductSeasonality_ProductID
        ON dbo.ProductSeasonality(ProductID);

    CREATE INDEX IX_ProductSeasonality_MonthWindow
        ON dbo.ProductSeasonality(StartMonth, EndMonth, IsYearRound, HasPeakSeason)
        INCLUDE (ProductID, SeasonType, SeasonLabel, PeakMonths, IsOffSeason);
END;
GO

IF COL_LENGTH(N'dbo.ProductSeasonality', N'HasPeakSeason') IS NULL
BEGIN
    ALTER TABLE dbo.ProductSeasonality
        ADD HasPeakSeason bit NOT NULL
            CONSTRAINT DF_ProductSeasonality_HasPeakSeason DEFAULT (0);
END;
GO

IF COL_LENGTH(N'dbo.ProductSeasonality', N'IsControlledCultivation') IS NULL
BEGIN
    ALTER TABLE dbo.ProductSeasonality
        ADD IsControlledCultivation bit NOT NULL
            CONSTRAINT DF_ProductSeasonality_IsControlledCultivation DEFAULT (0);
END;
GO

IF COL_LENGTH(N'dbo.ProductSeasonality', N'IsImportedSeason') IS NULL
BEGIN
    ALTER TABLE dbo.ProductSeasonality
        ADD IsImportedSeason bit NOT NULL
            CONSTRAINT DF_ProductSeasonality_IsImportedSeason DEFAULT (0);
END;
GO

IF COL_LENGTH(N'dbo.ProductSeasonality', N'IsOffSeason') IS NULL
BEGIN
    ALTER TABLE dbo.ProductSeasonality
        ADD IsOffSeason bit NOT NULL
            CONSTRAINT DF_ProductSeasonality_IsOffSeason DEFAULT (0);
END;
GO

IF COL_LENGTH(N'dbo.ProductSeasonality', N'IsPostHarvestAvailability') IS NULL
BEGIN
    ALTER TABLE dbo.ProductSeasonality
        ADD IsPostHarvestAvailability bit NOT NULL
            CONSTRAINT DF_ProductSeasonality_IsPostHarvestAvailability DEFAULT (0);
END;
GO

IF COL_LENGTH(N'dbo.ProductSeasonality', N'SeasonScoreWeight') IS NULL
BEGIN
    ALTER TABLE dbo.ProductSeasonality
        ADD SeasonScoreWeight int NOT NULL
            CONSTRAINT DF_ProductSeasonality_SeasonScoreWeight DEFAULT (50);
END;
GO

IF COL_LENGTH(N'dbo.ProductSeasonality', N'ConfidenceLevel') IS NULL
BEGIN
    ALTER TABLE dbo.ProductSeasonality
        ADD ConfidenceLevel nvarchar(40) NULL;
END;
GO

IF COL_LENGTH(N'dbo.ProductSeasonality', N'ProductName') IS NOT NULL
BEGIN
    ALTER TABLE dbo.ProductSeasonality DROP COLUMN ProductName;
END;
GO

IF COL_LENGTH(N'dbo.ProductSeasonality', N'CategoryName') IS NOT NULL
BEGIN
    ALTER TABLE dbo.ProductSeasonality DROP COLUMN CategoryName;
END;
GO

IF COL_LENGTH(N'dbo.ProductSeasonality', N'DbRuleHint') IS NOT NULL
BEGIN
    ALTER TABLE dbo.ProductSeasonality DROP COLUMN DbRuleHint;
END;
GO

IF COL_LENGTH(N'dbo.ProductSeasonality', N'ReviewNote') IS NOT NULL
BEGIN
    ALTER TABLE dbo.ProductSeasonality DROP COLUMN ReviewNote;
END;
GO

;WITH SourceRows
(
    ProductID,
    Country,
    ProvinceRegion,
    AreaDetail,
    SeasonType,
    SeasonLabel,
    StartMonth,
    EndMonth,
    PeakMonths,
    IsYearRound,
    HasPeakSeason,
    IsControlledCultivation,
    IsImportedSeason,
    IsOffSeason,
    IsPostHarvestAvailability,
    SeasonScoreWeight,
    ConfidenceLevel,
    SourceNote
)
AS
(
    SELECT *
    FROM
    (
        VALUES
{values_sql}
    ) AS source
    (
        ProductID,
        Country,
        ProvinceRegion,
        AreaDetail,
        SeasonType,
        SeasonLabel,
        StartMonth,
        EndMonth,
        PeakMonths,
        IsYearRound,
        HasPeakSeason,
        IsControlledCultivation,
        IsImportedSeason,
        IsOffSeason,
        IsPostHarvestAvailability,
        SeasonScoreWeight,
        ConfidenceLevel,
        SourceNote
    )
),
ValidRows AS
(
    SELECT source.*
    FROM SourceRows AS source
    INNER JOIN dbo.Products AS product
        ON product.ProductID = source.ProductID
)
MERGE dbo.ProductSeasonality AS target
USING ValidRows AS source
ON target.ProductID = source.ProductID
AND target.SeasonLabel = source.SeasonLabel
AND target.StartMonth = source.StartMonth
AND target.EndMonth = source.EndMonth
WHEN MATCHED AND
(
    ISNULL(target.Country, N'') <> ISNULL(source.Country, N'')
    OR ISNULL(target.ProvinceRegion, N'') <> ISNULL(source.ProvinceRegion, N'')
    OR ISNULL(target.AreaDetail, N'') <> ISNULL(source.AreaDetail, N'')
    OR ISNULL(target.SeasonType, N'') <> ISNULL(source.SeasonType, N'')
    OR ISNULL(target.PeakMonths, N'') <> ISNULL(source.PeakMonths, N'')
    OR ISNULL(target.IsYearRound, 0) <> ISNULL(source.IsYearRound, 0)
    OR ISNULL(target.HasPeakSeason, 0) <> ISNULL(source.HasPeakSeason, 0)
    OR ISNULL(target.IsControlledCultivation, 0) <> ISNULL(source.IsControlledCultivation, 0)
    OR ISNULL(target.IsImportedSeason, 0) <> ISNULL(source.IsImportedSeason, 0)
    OR ISNULL(target.IsOffSeason, 0) <> ISNULL(source.IsOffSeason, 0)
    OR ISNULL(target.IsPostHarvestAvailability, 0) <> ISNULL(source.IsPostHarvestAvailability, 0)
    OR ISNULL(target.SeasonScoreWeight, 0) <> ISNULL(source.SeasonScoreWeight, 0)
    OR ISNULL(target.ConfidenceLevel, N'') <> ISNULL(source.ConfidenceLevel, N'')
    OR ISNULL(target.SourceNote, N'') <> ISNULL(source.SourceNote, N'')
)
THEN UPDATE SET
    Country = source.Country,
    ProvinceRegion = source.ProvinceRegion,
    AreaDetail = source.AreaDetail,
    SeasonType = source.SeasonType,
    PeakMonths = source.PeakMonths,
    IsYearRound = source.IsYearRound,
    HasPeakSeason = source.HasPeakSeason,
    IsControlledCultivation = source.IsControlledCultivation,
    IsImportedSeason = source.IsImportedSeason,
    IsOffSeason = source.IsOffSeason,
    IsPostHarvestAvailability = source.IsPostHarvestAvailability,
    SeasonScoreWeight = source.SeasonScoreWeight,
    ConfidenceLevel = source.ConfidenceLevel,
    SourceNote = source.SourceNote,
    UpdatedAt = SYSUTCDATETIME()
WHEN NOT MATCHED BY TARGET
THEN INSERT
(
    ProductID,
    Country,
    ProvinceRegion,
    AreaDetail,
    SeasonType,
    SeasonLabel,
    StartMonth,
    EndMonth,
    PeakMonths,
    IsYearRound,
    HasPeakSeason,
    IsControlledCultivation,
    IsImportedSeason,
    IsOffSeason,
    IsPostHarvestAvailability,
    SeasonScoreWeight,
    ConfidenceLevel,
    SourceNote
)
VALUES
(
    source.ProductID,
    source.Country,
    source.ProvinceRegion,
    source.AreaDetail,
    source.SeasonType,
    source.SeasonLabel,
    source.StartMonth,
    source.EndMonth,
    source.PeakMonths,
    source.IsYearRound,
    source.HasPeakSeason,
    source.IsControlledCultivation,
    source.IsImportedSeason,
    source.IsOffSeason,
    source.IsPostHarvestAvailability,
    source.SeasonScoreWeight,
    source.ConfidenceLevel,
    source.SourceNote
);
GO

DECLARE @ExpectedRows int =
(
    SELECT COALESCE(SUM(expected.SeasonRows), 0)
    FROM
    (
        VALUES
{expected_counts_sql}
    ) AS expected(ProductID, SeasonRows)
    INNER JOIN dbo.Products AS product
        ON product.ProductID = expected.ProductID
);
DECLARE @ActualRows int =
(
    SELECT COUNT(*)
    FROM dbo.ProductSeasonality
    WHERE ProductID IN (SELECT DISTINCT ProductID FROM dbo.Products)
);

IF @ActualRows < @ExpectedRows
BEGIN
    RAISERROR(N'ProductSeasonality seed thieu dong. Expected at least %d, actual %d.', 16, 1, @ExpectedRows, @ActualRows);
    RETURN;
END;
GO
"""


def build_verify(rows: list[dict[str, object]]) -> str:
    expected_counts_sql = build_expected_count_values(rows)
    return f"""SET NOCOUNT ON;

IF OBJECT_ID(N'dbo.ProductSeasonality', N'U') IS NULL
BEGIN
    RAISERROR(N'ProductSeasonality chua ton tai.', 16, 1);
    RETURN;
END;

SELECT
    TotalSeasonalityRows = COUNT(*),
    DistinctProducts = COUNT(DISTINCT ProductID),
    YearRoundRows = SUM(CASE WHEN IsYearRound = 1 THEN 1 ELSE 0 END),
    SeasonalRows = SUM(CASE WHEN IsYearRound = 0 THEN 1 ELSE 0 END),
    PeakSeasonRows = SUM(CASE WHEN HasPeakSeason = 1 THEN 1 ELSE 0 END),
    ControlledCultivationRows = SUM(CASE WHEN IsControlledCultivation = 1 THEN 1 ELSE 0 END),
    ImportedSeasonRows = SUM(CASE WHEN IsImportedSeason = 1 THEN 1 ELSE 0 END),
    OffSeasonRows = SUM(CASE WHEN IsOffSeason = 1 THEN 1 ELSE 0 END),
    PostHarvestRows = SUM(CASE WHEN IsPostHarvestAvailability = 1 THEN 1 ELSE 0 END),
    MinSeasonScoreWeight = MIN(SeasonScoreWeight),
    MaxSeasonScoreWeight = MAX(SeasonScoreWeight),
    ConfidenceLevelFilledRows = SUM(CASE WHEN ConfidenceLevel IS NOT NULL THEN 1 ELSE 0 END)
FROM dbo.ProductSeasonality;

DECLARE @ExpectedRows int =
(
    SELECT COALESCE(SUM(expected.SeasonRows), 0)
    FROM
    (
        VALUES
{expected_counts_sql}
    ) AS expected(ProductID, SeasonRows)
    INNER JOIN dbo.Products AS product
        ON product.ProductID = expected.ProductID
);

DECLARE @ExpectedProducts int =
(
    SELECT COUNT(*)
    FROM
    (
        VALUES
{expected_counts_sql}
    ) AS expected(ProductID, SeasonRows)
    INNER JOIN dbo.Products AS product
        ON product.ProductID = expected.ProductID
);

SELECT
    season.ProductID,
    ProductName = MAX(product.ProductName),
    SeasonRows = COUNT(*)
FROM dbo.ProductSeasonality AS season
INNER JOIN dbo.Products AS product
    ON product.ProductID = season.ProductID
GROUP BY season.ProductID
HAVING COUNT(*) > 1
ORDER BY season.ProductID;

SELECT
    MissingProductRows = COUNT(*)
FROM dbo.ProductSeasonality AS season
LEFT JOIN dbo.Products AS product
    ON product.ProductID = season.ProductID
WHERE product.ProductID IS NULL;

IF
(
    SELECT COUNT(*)
    FROM dbo.ProductSeasonality
) < @ExpectedRows
BEGIN
    RAISERROR(N'So dong ProductSeasonality thap hon seed hop le theo Products hien co.', 16, 1);
    RETURN;
END;

IF
(
    SELECT COUNT(DISTINCT ProductID)
    FROM dbo.ProductSeasonality
) < @ExpectedProducts
BEGIN
    RAISERROR(N'So san pham co mua vu thap hon seed hop le theo Products hien co.', 16, 1);
    RETURN;
END;

IF EXISTS
(
    SELECT 1
    FROM dbo.ProductSeasonality
    WHERE StartMonth NOT BETWEEN 1 AND 12
       OR EndMonth NOT BETWEEN 1 AND 12
       OR SeasonScoreWeight NOT BETWEEN 0 AND 100
)
BEGIN
    RAISERROR(N'Co dong mua vu co StartMonth/EndMonth hoac SeasonScoreWeight ngoai khoang hop le.', 16, 1);
    RETURN;
END;
"""


def main() -> None:
    rows = read_rows()
    if not rows:
        raise RuntimeError("No seasonality rows found.")
    OUTPUT_DELTA.write_text(build_delta(rows), encoding="utf-8-sig")
    OUTPUT_VERIFY.write_text(build_verify(rows), encoding="utf-8-sig")
    print(f"Wrote {OUTPUT_DELTA}")
    print(f"Wrote {OUTPUT_VERIFY}")
    print(f"Rows: {len(rows)}")
    print(f"Products: {len({int(row['ProductID']) for row in rows})}")


if __name__ == "__main__":
    main()

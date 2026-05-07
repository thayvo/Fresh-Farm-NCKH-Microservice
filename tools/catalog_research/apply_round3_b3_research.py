from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK_PATH = REPO_ROOT / "output" / "spreadsheet" / "recommendation_product_data_review.xlsx"
ROUND3_COLLECTION_TIMESTAMP = "2026-04-24 21:05:08 ICT"


SEASONALITY_ROWS = [
    (
        "B3",
        65,
        "Dưa hấu không hạt",
        "Trái cây",
        "Long An, Việt Nam",
        "District",
        "Vĩnh Hưng, Long An",
        "Dưa hấu Tết",
        12,
        1,
        "12-1",
        "Yes",
        "Báo Long An ghi nhận nông dân Vĩnh Hưng gieo trồng hơn 250ha dưa hấu phục vụ thị trường Tết, dự kiến thu hoạch vào giữa tháng Chạp.",
        "Vĩnh Hưng: Nông dân gieo trồng hơn 250ha dưa hấu phục vụ thị trường tết",
        "Báo Long An",
        "https://baolongan.vn/vinh-hung-nong-dan-gieo-trong-hon-250ha-dua-hau-phuc-vu-thi-truong-tet-a186458.html",
        "2024-12-01",
        "Medium",
        ROUND3_COLLECTION_TIMESTAMP,
        "Nguồn xác nhận vụ Tết tại Long An nhưng không tách riêng giống không hạt.",
        "Dùng chung cho SKU 65 và 118 đến khi có nguồn riêng cho giống không hạt/mini.",
    ),
    (
        "B3",
        66,
        "Bơ Booth 7",
        "Trái cây",
        "Đắk Lắk, Việt Nam",
        "Province",
        "Đắk Lắk",
        "Chính vụ và vụ muộn",
        4,
        11,
        "4-7; 9-11",
        "Yes",
        "Sở Công Thương Đắk Lắk ghi cây bơ có chính vụ từ tháng 4-7 và vụ muộn từ tháng 9-11.",
        "Bơ Đắk Lắk - Hành trình đến với người tiêu dùng",
        "Sở Công Thương Đắk Lắk",
        "https://socongthuong.daklak.gov.vn/vi/news/hoat-dong-nganh-cong-thuong/bo-dak-lak-hanh-trinh-den-voi-nguoi-tieu-dung-3497.html",
        "",
        "High",
        ROUND3_COLLECTION_TIMESTAMP,
        "Nguồn nói bơ Đắk Lắk nói chung, trong đó có Booth 7; chưa tách riêng tuyệt đối lịch Booth 7 theo từng vùng nhỏ.",
        "",
    ),
    (
        "B3",
        67,
        "Dâu tây Đà Lạt",
        "Trái cây",
        "Đà Lạt, Việt Nam",
        "City / nearby district",
        "Đà Lạt - Lạc Dương, Lâm Đồng",
        "Chu kỳ thu hoạch vụ nhà kính",
        "Sau 4 tháng xuống giống",
        "khoảng 6 tháng sau đó",
        "Mùa nắng cho sản lượng tốt hơn",
        "No",
        "Nguồn Lâm Đồng ghi giống PS8.07/PS8.10 sau khoảng 4 tháng xuống giống bước vào thu hoạch liên tục đến khoảng 6 tháng sau đó; mùa mưa làm sản lượng giảm mạnh so với mùa nắng.",
        "Khảo nghiệm 2 giống dâu tây mới ở Đà Lạt",
        "Thị trấn Lạc Dương, Lâm Đồng",
        "https://thitranlacduong.lamdong.gov.vn/chi-tiet-tin-tuc/?param=khao-nghiem-2-giong-dau-tay-moi-o-da-lat",
        "",
        "High",
        ROUND3_COLLECTION_TIMESTAMP,
        "Dâu tây công nghệ cao có thể trồng/thu quanh năm theo lứa; không nên biến thành tháng cố định nếu không có ngày xuống giống.",
        "Dùng thêm nguồn Khuyến nông Lâm Đồng về mùa mưa làm giảm sản lượng.",
    ),
    (
        "B3",
        72,
        "Mít Thái siêu ngọt",
        "Trái cây",
        "Tiền Giang, Việt Nam",
        "Province / specialized areas",
        "Cái Bè, Cai Lậy, Châu Thành, Tân Phước, TX Cai Lậy",
        "Rải vụ gần như quanh năm",
        1,
        12,
        "Year-round",
        "No",
        "Cổng TTĐT Tiền Giang ghi mít Thái siêu sớm cho thu hoạch rải vụ gần như quanh năm.",
        "Giá mít Thái tăng mạnh, giúp nông dân ổn định cuộc sống trong tình hình dịch bệnh còn phức tạp",
        "Cổng thông tin điện tử tỉnh Tiền Giang",
        "https://tiengiang.gov.vn/chi-tiet-tin?%2Fgia-mit-thai-tang-manh-giup-nong-dan-on-inh-cuoc-song-trong-tinh-hinh-dich-benh-con-phuc-tap%2F32403960=",
        "",
        "High",
        ROUND3_COLLECTION_TIMESTAMP,
        "",
        "",
    ),
    (
        "B3",
        116,
        "Chuối già Việt Nam",
        "Trái cây",
        "Tiền Giang, Việt Nam",
        "Domestic tropical crop",
        "Việt Nam / Tiền Giang chưa có nguồn tỉnh đủ chi tiết",
        "Thu hoạch quanh năm - bối cảnh cây chuối",
        1,
        12,
        "1-4 sản lượng nhiều theo nguồn khuyến nông",
        "No",
        "Nguồn Khuyến nông Việt Nam ghi cây chuối có thể thu hoạch quanh năm, sản lượng nhiều từ tháng 1 đến tháng 4; nguồn không riêng Tiền Giang.",
        "Đồng Tháp: Giá chuối tăng, nông dân vẫn lo",
        "Trung tâm Khuyến nông Quốc gia",
        "https://khuyennongvn.gov.vn/thong-tin-thi-truong/thi-truong-trong-nuoc/dong-thap-gia-chuoi-tang-nong-dan-van-lo-14376.html",
        "",
        "Medium",
        ROUND3_COLLECTION_TIMESTAMP,
        "Chưa tìm được nguồn cực mạnh riêng cho chuối già Tiền Giang; tạm dùng nguồn khuyến nông cho cây chuối vùng tương đồng và phải kiểm tiếp.",
        "",
    ),
    (
        "B3",
        118,
        "Dưa hấu không hạt mini",
        "Trái cây",
        "Long An, Việt Nam",
        "District",
        "Vĩnh Hưng, Long An",
        "Dưa hấu Tết",
        12,
        1,
        "12-1",
        "Yes",
        "Báo Long An ghi nhận nông dân Vĩnh Hưng gieo trồng hơn 250ha dưa hấu phục vụ thị trường Tết, dự kiến thu hoạch vào giữa tháng Chạp.",
        "Vĩnh Hưng: Nông dân gieo trồng hơn 250ha dưa hấu phục vụ thị trường tết",
        "Báo Long An",
        "https://baolongan.vn/vinh-hung-nong-dan-gieo-trong-hon-250ha-dua-hau-phuc-vu-thi-truong-tet-a186458.html",
        "2024-12-01",
        "Medium",
        ROUND3_COLLECTION_TIMESTAMP,
        "Nguồn xác nhận vụ Tết tại Long An nhưng không tách riêng giống không hạt mini.",
        "",
    ),
    (
        "B3",
        119,
        "Thanh long ruột trắng",
        "Trái cây",
        "Bình Thuận, Việt Nam",
        "Province / GI area",
        "Bình Thuận",
        "Vụ nghịch cuối năm - đầu năm",
        "10 âm lịch",
        "sau Tết",
        "tháng Chạp - sau Tết",
        "No",
        "Sở NN&PTNT Bình Thuận ghi nông dân chong đèn từ tháng 10 âm lịch và các lứa nghịch vụ thu hoạch từ tháng Chạp trở đi.",
        "Thanh long vào vụ nghịch: [Bài 1] Nâng cao chất lượng đáp ứng thị trường",
        "Sở NN&PTNT Bình Thuận",
        "https://snnptnt.binhthuan.gov.vn/tin-nong-nghiep/thanh-long-vao-vu-nghich-bai-1-nang-cao-chat-luong-dap-ung-thi-truong-867366",
        "2023-12-05",
        "High",
        ROUND3_COLLECTION_TIMESTAMP,
        "Nguồn ghi thanh long Bình Thuận nói chung, chưa tách riêng ruột trắng; mốc tháng ở dạng âm lịch nên chưa đổi sang dương lịch.",
        "",
    ),
    (
        "B3",
        120,
        "Sầu riêng Ri6",
        "Trái cây",
        "Đắk Lắk, Việt Nam",
        "Province",
        "Đắk Lắk",
        "Vụ thu hoạch chính ở Đắk Lắk",
        8,
        9,
        "8-9",
        "Yes",
        "Khuyến nông Việt Nam ghi sầu riêng Đắk Lắk bắt đầu vào vụ; chính vụ giống Dona tại Krông Pắc từ 27/8 đến đầu tháng 9/2022; Ri6 được nêu cùng nhóm giống chất lượng trong nguồn khác của địa phương.",
        "Đắk Lắk: Sầu riêng được giá, nhưng năng suất năm nay lại thấp",
        "Trung tâm Khuyến nông Quốc gia",
        "https://khuyennongvn.gov.vn/chuong-trinh-nganh-nong-nghiep/tai-co-cau-nganh-nong-nghiep/dak-lak-sau-rieng-duoc-gia-nhung-nang-suat-nam-nay-lai-thap-22052.html",
        "",
        "Medium",
        ROUND3_COLLECTION_TIMESTAMP,
        "Nguồn mùa vụ mạnh cho sầu riêng Đắk Lắk nhưng không tách riêng Ri6; cần thêm nguồn cultivar Ri6 tại Đắk Lắk nếu dùng làm claim chính xác tuyệt đối.",
        "",
    ),
]


RESEARCH_ROWS = [
    (
        "B3",
        65,
        "Dưa hấu không hạt",
        "Origin/locality context",
        "Long An có vùng dưa hấu Vĩnh Hưng phục vụ thị trường Tết; dưa Long Trì, Tân Trụ là các tên được địa phương ghi nhận.",
        "Long An",
        "Góc ảnh 'Long An quê hương tôi': Dưa hấu Long An",
        "Báo Long An",
        "https://baolongan.vn/goc-anh-long-an-que-huong-toi-dua-hau-long-an-a180898.html",
        "2024-08-16",
        "Medium",
        ROUND3_COLLECTION_TIMESTAMP,
        "Báo địa phương uy tín nhưng chưa phải hồ sơ kỹ thuật/cây trồng cấp sở.",
        "",
    ),
    (
        "B3",
        66,
        "Bơ Booth 7",
        "Origin/production context",
        "Đắk Lắk có các giống bơ phổ biến gồm Booth 7, bơ tứ quý, bơ 034; diện tích cho sản phẩm 7.228 ha năm 2021 theo Sở NN&PTNT được Sở Công Thương dẫn lại.",
        "Đắk Lắk",
        "Bơ Đắk Lắk - Hành trình đến với người tiêu dùng",
        "Sở Công Thương Đắk Lắk",
        "https://socongthuong.daklak.gov.vn/vi/news/hoat-dong-nganh-cong-thuong/bo-dak-lak-hanh-trinh-den-voi-nguoi-tieu-dung-3497.html",
        "",
        "High",
        ROUND3_COLLECTION_TIMESTAMP,
        "",
        "",
    ),
    (
        "B3",
        67,
        "Dâu tây Đà Lạt",
        "Preservation guidance",
        "Trái dâu tây không bảo quản được lâu, nên bảo quản và vận chuyển trong điều kiện lạnh sau thu hoạch.",
        "Đà Lạt/Lâm Đồng",
        "Quy trình kỹ thuật trồng cây dâu tây",
        "Khuyến nông Lâm Đồng",
        "https://khuyennong.lamdong.gov.vn/ky-thuat-trong-trot/ki-thuat-trong-rau/865-quy-trinh-k-thut-trng-cay-dau-tay",
        "",
        "High",
        ROUND3_COLLECTION_TIMESTAMP,
        "",
        "",
    ),
    (
        "B3",
        67,
        "Dâu tây Đà Lạt",
        "Flavor / texture context",
        "Các giống tại Đà Lạt/Lạc Dương có mô tả như Mỹ Hương ngọt thơm, New Zealand giòn ngọt thơm, Akihime mềm thơm.",
        "Đà Lạt/Lâm Đồng",
        "Dâu tây Đà Lạt với những biện pháp quản lý dịch hại mới",
        "Khuyến nông Lâm Đồng",
        "https://khuyennong.lamdong.gov.vn/tin-tuc-su-kien/1518-dau-tay-da-lat-voi-nhung-bien-phap-quan-ly-dich-hai-moi",
        "",
        "High",
        ROUND3_COLLECTION_TIMESTAMP,
        "Flavor/texture phụ thuộc giống cụ thể; catalog chỉ ghi Dâu tây Đà Lạt nên cần giữ ở mức tổng quát.",
        "",
    ),
    (
        "B3",
        72,
        "Mít Thái siêu ngọt",
        "Origin/production context",
        "Tiền Giang có vùng mít Thái chuyên canh tại Cái Bè, Cai Lậy, Châu Thành, Tân Phước và TX Cai Lậy; gần 13.000 ha đang cho thu hoạch.",
        "Tiền Giang",
        "Phát triển vùng mít chuyên canh trên những địa bàn sản xuất khó khăn",
        "Cổng thông tin điện tử tỉnh Tiền Giang",
        "https://tiengiang.gov.vn/chi-tiet-tin/?%2Fphat-trien-vung-mit-chuyen-canh-tren-nhung-ia-ban-san-xuat-kho-khan%2F58616436=",
        "",
        "High",
        ROUND3_COLLECTION_TIMESTAMP,
        "",
        "",
    ),
    (
        "B3",
        117,
        "Ổi xá lị",
        "Origin/locality status",
        "Catalog ghi Long An nhưng chưa tìm được nguồn kỹ thuật/nguồn nhà nước đủ mạnh riêng cho ổi xá lị Long An.",
        "Long An",
        "Lan tỏa phong trào khởi nghiệp",
        "Báo Long An",
        "https://baolongan.vn/lan-toa-phong-trao-khoi-nghiep-a41165.html",
        "",
        "Low",
        ROUND3_COLLECTION_TIMESTAMP,
        "Chỉ xác nhận mô hình trồng ổi tại Long An ở mức báo địa phương; chưa đủ để chốt mùa vụ.",
        "",
    ),
    (
        "B3",
        120,
        "Sầu riêng Ri6",
        "Origin/brand context",
        "Đắk Lắk có nhãn hiệu 'Sầu riêng Krông Pắc' và 'Sầu riêng Cư M’gar'; Ri6 được nhắc cùng nhóm giống chất lượng như Dona, Cái Mơn.",
        "Đắk Lắk",
        "Tỉnh Đắk Lắk tổ chức hội nghị tổng kết ngành hàng sầu riêng năm 2023, phương hướng nhiệm vụ năm 2024",
        "Sở Công Thương Đắk Lắk",
        "https://socongthuong.daklak.gov.vn/vi/news/hoat-dong-nganh-cong-thuong/to-chuc-hoi-nghi-tong-ket-nganh-hang-sau-rieng-nam-2023-phuong-huong-nhiem-vu-nam-2024-4961.html",
        "",
        "High",
        ROUND3_COLLECTION_TIMESTAMP,
        "Nguồn xác nhận bối cảnh ngành hàng/nhãn hiệu địa phương; mùa vụ riêng Ri6 vẫn cần nguồn cultivar cụ thể.",
        "",
    ),
]


PRODUCT_SUMMARIES = {
    65: {
        "Review Status": "Partially verified - batch 3",
        "Verified Seasonality Summary": "Long An: có vụ dưa hấu Tết tại Vĩnh Hưng, dự kiến thu hoạch giữa tháng Chạp; chưa có nguồn tách riêng dưa không hạt.",
        "Verified Origin Country": "Vietnam",
        "Verified Origin Province/Region": "Long An - Vĩnh Hưng/Tân Trụ/Long Trì context",
        "Verified Domestic/Imported": "Domestic",
        "Primary Source URL": "https://baolongan.vn/vinh-hung-nong-dan-gieo-trong-hon-250ha-dua-hau-phuc-vu-thi-truong-tet-a186458.html",
        "Secondary Source URL": "https://baolongan.vn/goc-anh-long-an-que-huong-toi-dua-hau-long-an-a180898.html",
        "Thời gian thu thập": ROUND3_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "Nguồn xác nhận dưa hấu Long An/vụ Tết nhưng chưa tách riêng giống không hạt.",
    },
    66: {
        "Review Status": "Partially verified - batch 3",
        "Verified Seasonality Summary": "Đắk Lắk: bơ có chính vụ tháng 4-7 và vụ muộn tháng 9-11; nguồn nêu Booth 7 là một giống phổ biến.",
        "Verified Origin Country": "Vietnam",
        "Verified Origin Province/Region": "Đắk Lắk",
        "Verified Domestic/Imported": "Domestic",
        "Primary Source URL": "https://socongthuong.daklak.gov.vn/vi/news/hoat-dong-nganh-cong-thuong/bo-dak-lak-hanh-trinh-den-voi-nguoi-tieu-dung-3497.html",
        "Secondary Source URL": "https://khuyennongvn.gov.vn/guong-san-xuat-gioi/kinh-nghiem-su-dung-dam-ca-trong-canh-tac-bo-cua-mot-nong-dan-dak-nong-22533.html",
        "Thời gian thu thập": ROUND3_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "Mùa vụ là cho bơ Đắk Lắk nói chung; chưa có lịch Booth 7 riêng theo từng huyện.",
    },
    67: {
        "Review Status": "Partially verified - batch 3",
        "Verified Seasonality Summary": "Đà Lạt/Lạc Dương: dâu tây công nghệ cao thu theo lứa, sau khoảng 4 tháng xuống giống có thể thu liên tục khoảng 6 tháng; mùa mưa làm giảm sản lượng so với mùa nắng.",
        "Verified Origin Country": "Vietnam",
        "Verified Origin Province/Region": "Lâm Đồng - Đà Lạt/Lạc Dương",
        "Verified Domestic/Imported": "Domestic",
        "Verified Preservation Guidance": "Bảo quản và vận chuyển lạnh sau thu hoạch.",
        "Verified Flavor Profile": "Ngọt thơm, tùy giống",
        "Verified Texture Profile": "Mềm/giòn tùy giống",
        "Primary Source URL": "https://thitranlacduong.lamdong.gov.vn/chi-tiet-tin-tuc/?param=khao-nghiem-2-giong-dau-tay-moi-o-da-lat",
        "Secondary Source URL": "https://khuyennong.lamdong.gov.vn/ky-thuat-trong-trot/ki-thuat-trong-rau/865-quy-trinh-k-thut-trng-cay-dau-tay",
        "Thời gian thu thập": ROUND3_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "Không nên quy về tháng cố định nếu không biết ngày xuống giống/giống/nhà kính; flavor phụ thuộc giống.",
    },
    72: {
        "Review Status": "Partially verified - batch 3",
        "Verified Seasonality Summary": "Tiền Giang: mít Thái siêu sớm cho thu hoạch rải vụ gần như quanh năm.",
        "Verified Origin Country": "Vietnam",
        "Verified Origin Province/Region": "Tiền Giang - Cái Bè/Cai Lậy/Châu Thành/Tân Phước/TX Cai Lậy",
        "Verified Domestic/Imported": "Domestic",
        "Primary Source URL": "https://tiengiang.gov.vn/chi-tiet-tin?%2Fgia-mit-thai-tang-manh-giup-nong-dan-on-inh-cuoc-song-trong-tinh-hinh-dich-benh-con-phuc-tap%2F32403960=",
        "Secondary Source URL": "https://tiengiang.gov.vn/chi-tiet-tin/?%2Fphat-trien-vung-mit-chuyen-canh-tren-nhung-ia-ban-san-xuat-kho-khan%2F58616436=",
        "Thời gian thu thập": ROUND3_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "",
    },
    116: {
        "Review Status": "Partially verified - batch 3",
        "Verified Seasonality Summary": "Chuối có thể thu hoạch quanh năm, sản lượng nhiều từ tháng 1-4 theo nguồn khuyến nông; chưa có nguồn riêng đủ mạnh cho chuối già Tiền Giang.",
        "Verified Origin Country": "Vietnam",
        "Verified Origin Province/Region": "Tiền Giang (catalog); season source from regional banana context",
        "Verified Domestic/Imported": "Domestic",
        "Primary Source URL": "https://khuyennongvn.gov.vn/thong-tin-thi-truong/thi-truong-trong-nuoc/dong-thap-gia-chuoi-tang-nong-dan-van-lo-14376.html",
        "Secondary Source URL": "",
        "Thời gian thu thập": ROUND3_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "Nguồn mùa vụ không riêng Tiền Giang và không tách riêng chuối già; cần kiểm thêm nếu dùng cho rule chính thức.",
    },
    117: {
        "Review Status": "Partially verified - batch 3",
        "Verified Seasonality Summary": "Chưa có mùa vụ đủ mạnh cho ổi xá lị Long An; mới xác nhận bối cảnh có mô hình trồng ổi tại Long An ở mức báo địa phương.",
        "Verified Origin Country": "Vietnam",
        "Verified Origin Province/Region": "Long An",
        "Verified Domestic/Imported": "Domestic",
        "Primary Source URL": "https://baolongan.vn/lan-toa-phong-trao-khoi-nghiep-a41165.html",
        "Secondary Source URL": "",
        "Thời gian thu thập": ROUND3_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "Chưa đủ nguồn để chốt mùa vụ/đặc tính riêng ổi xá lị Long An.",
    },
    118: {
        "Review Status": "Partially verified - batch 3",
        "Verified Seasonality Summary": "Long An: có vụ dưa hấu Tết tại Vĩnh Hưng, dự kiến thu hoạch giữa tháng Chạp; chưa có nguồn tách riêng giống không hạt mini.",
        "Verified Origin Country": "Vietnam",
        "Verified Origin Province/Region": "Long An - Vĩnh Hưng/Tân Trụ/Long Trì context",
        "Verified Domestic/Imported": "Domestic",
        "Primary Source URL": "https://baolongan.vn/vinh-hung-nong-dan-gieo-trong-hon-250ha-dua-hau-phuc-vu-thi-truong-tet-a186458.html",
        "Secondary Source URL": "https://baolongan.vn/goc-anh-long-an-que-huong-toi-dua-hau-long-an-a180898.html",
        "Thời gian thu thập": ROUND3_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "Nguồn xác nhận dưa hấu Long An/vụ Tết nhưng chưa tách riêng giống không hạt mini.",
    },
    119: {
        "Review Status": "Partially verified - batch 3",
        "Verified Seasonality Summary": "Bình Thuận: vụ nghịch bắt đầu chong đèn từ tháng 10 âm lịch, thu hoạch từ tháng Chạp trở đi; dùng chung nguồn thanh long Bình Thuận.",
        "Verified Origin Country": "Vietnam",
        "Verified Origin Province/Region": "Bình Thuận - vùng GI thanh long",
        "Verified Domestic/Imported": "Domestic",
        "Primary Source URL": "https://snnptnt.binhthuan.gov.vn/tin-nong-nghiep/thanh-long-vao-vu-nghich-bai-1-nang-cao-chat-luong-dap-ung-thi-truong-867366",
        "Secondary Source URL": "https://ipvietnam.gov.vn/phat-trien-chi-dan-ia-ly/-/asset_publisher/SGA9PgvmYtWI/content/thanh-long-binh-thuan-duoc-bao-ho-cddl-tai-nhat-ban?inheritRedirect=false",
        "Thời gian thu thập": ROUND3_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "Nguồn chưa tách riêng ruột trắng; mốc mùa vụ theo âm lịch.",
    },
    120: {
        "Review Status": "Partially verified - batch 3",
        "Verified Seasonality Summary": "Đắk Lắk: sầu riêng vào vụ khoảng tháng 8-9 theo nguồn khuyến nông; chưa có lịch riêng Ri6 tại Đắk Lắk.",
        "Verified Origin Country": "Vietnam",
        "Verified Origin Province/Region": "Đắk Lắk - Krông Pắc/Cư M'gar context",
        "Verified Domestic/Imported": "Domestic",
        "Primary Source URL": "https://khuyennongvn.gov.vn/chuong-trinh-nganh-nong-nghiep/tai-co-cau-nganh-nong-nghiep/dak-lak-sau-rieng-duoc-gia-nhung-nang-suat-nam-nay-lai-thap-22052.html",
        "Secondary Source URL": "https://socongthuong.daklak.gov.vn/vi/news/hoat-dong-nganh-cong-thuong/to-chuc-hoi-nghi-tong-ket-nganh-hang-sau-rieng-nam-2023-phuong-huong-nhiem-vu-nam-2024-4961.html",
        "Thời gian thu thập": ROUND3_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "Nguồn mùa vụ là cho sầu riêng Đắk Lắk nói chung, chưa tách riêng cultivar Ri6.",
    },
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK_PATH))
    return parser.parse_args()


def normalize_row(row):
    return tuple(None if cell == "" else cell for cell in row)


def append_unique_rows(ws, rows):
    existing = {
        normalize_row(row)
        for row in ws.iter_rows(min_row=2, values_only=True)
        if any(cell is not None and cell != "" for cell in row)
    }
    for row in rows:
        normalized = normalize_row(row)
        if normalized not in existing:
            ws.append(normalized)
            existing.add(normalized)


def update_products_review(ws):
    headers = [cell.value for cell in ws[1]]
    header_index = {header: idx + 1 for idx, header in enumerate(headers)}
    product_row_map = {}
    for row in range(2, ws.max_row + 1):
        product_id = ws.cell(row=row, column=1).value
        if product_id is not None:
            product_row_map[product_id] = row

    for product_id, updates in PRODUCT_SUMMARIES.items():
        row = product_row_map.get(product_id)
        if row is None:
            continue
        for header, value in updates.items():
            ws.cell(row=row, column=header_index[header]).value = value


def main() -> None:
    args = parse_args()
    workbook_path = Path(args.workbook)
    wb = load_workbook(workbook_path)
    append_unique_rows(wb["Seasonality Evidence"], SEASONALITY_ROWS)
    append_unique_rows(wb["Research Evidence"], RESEARCH_ROWS)
    update_products_review(wb["Products Review"])
    wb.save(workbook_path)


if __name__ == "__main__":
    main()

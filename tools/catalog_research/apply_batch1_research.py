from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK_PATH = REPO_ROOT / "output" / "spreadsheet" / "recommendation_product_data_review_multibatch.xlsx"
BATCH1_COLLECTION_TIMESTAMP = "2026-04-24 14:43:30 ICT"


RAW_SEASONALITY_ROWS = [
    (
        "B1",
        61,
        "Táo Fuji Nhật",
        "Trái cây",
        "Nhật Bản",
        "Prefecture",
        "Aomori",
        "Fuji export / eating window",
        11,
        1,
        "11-12",
        "Yes",
        "Aomori export season starts in October; Fuji is highlighted in November and controlled storage extends quality beyond New Year.",
        "Fruit and Vegetables: Apples",
        "Japan Food Export Fair / Aomori Apple Council",
        "https://j-fec.com/en/producer/pdf/en/20/",
        "2022",
        "High",
        "End month is inferred from official storage note rather than an explicit Fuji close date.",
        "Use Aomori as the strongest verified locality currently available for Fuji export fruit.",
    ),
    (
        "B1",
        62,
        "Cam Cara ruột đỏ",
        "Trái cây",
        "Úc",
        "Growing regions",
        "Murray Valley / Riverina / Riverland",
        "Australia navel orange season",
        6,
        10,
        "6-8",
        "No",
        "Citrus Australia lists navel oranges in season June-October; Agriculture Victoria describes Cara Cara as an early-to-mid-season navel.",
        "What's in Season",
        "Citrus Australia",
        "https://citrusaustralia.com.au/buying-australian/season/",
        "",
        "High",
        "Cara Cara month range is inferred from the broader navel window because a month-by-month Cara Cara calendar was not found.",
        "Treat this as a provisional Australia-season mapping until a variety-specific official calendar is found.",
    ),
    (
        "B1",
        63,
        "Nho Mỹ không hạt",
        "Trái cây",
        "Mỹ",
        "Region",
        "Coachella Valley, California",
        "Early California table grape season",
        5,
        7,
        "5-6",
        "Yes",
        "California table grapes begin in Coachella Valley in May and usually finish in early July.",
        "Grapes From California Merchandising and Training Guide",
        "California Table Grape Commission",
        "https://www.grapesfromcalifornia.com/wp-content/uploads/2023/06/20230605-grapes-from-california-merchandising-and-training-guide-2023.pdf",
        "2023-06-05",
        "High",
        "",
        "",
    ),
    (
        "B1",
        63,
        "Nho Mỹ không hạt",
        "Trái cây",
        "Mỹ",
        "Region",
        "San Joaquin Valley, California",
        "Main California table grape season",
        7,
        12,
        "8-10",
        "Yes",
        "The San Joaquin Valley starts in early July and often continues into December.",
        "Grapes From California Merchandising and Training Guide",
        "California Table Grape Commission",
        "https://www.grapesfromcalifornia.com/wp-content/uploads/2023/06/20230605-grapes-from-california-merchandising-and-training-guide-2023.pdf",
        "2023-06-05",
        "High",
        "",
        "California grapes are marketed from May into January.",
    ),
    (
        "B1",
        71,
        "Măng cụt",
        "Trái cây",
        "Thái Lan",
        "Country",
        "Thailand",
        "National availability window",
        4,
        12,
        "4-8",
        "Yes",
        "The DOAE commodity page states availability from April to December with peak supply in April-August.",
        "Mangosteen",
        "Department of Agricultural Extension, Thailand",
        "https://www.doae.go.th/en/mangosteen-2/",
        "",
        "High",
        "",
        "",
    ),
    (
        "B1",
        71,
        "Măng cụt",
        "Trái cây",
        "Thái Lan",
        "Province / locality",
        "Khiriwong, Nakhon Si Thammarat",
        "Southern locality fruit season",
        7,
        9,
        "7-9",
        "No",
        "Thailand government tourism content notes fruit season in Khiriwong from July to September, including mangosteen.",
        "Kiriwong Village and Khao Luang Mountain",
        "Government of Thailand",
        "https://www.thailand.go.th/issue-focus-detail/001_02_255",
        "",
        "High",
        "",
        "Useful as a locality-specific southern Thailand season window.",
    ),
    (
        "B1",
        73,
        "Lựu đỏ Mỹ",
        "Trái cây",
        "Mỹ",
        "Region",
        "San Joaquin Valley, California",
        "California pomegranate season",
        10,
        1,
        "10-1",
        "Yes",
        "California pomegranates are described as being in season from October to January and are grown in the San Joaquin Valley.",
        "Why California",
        "POM Wonderful / California Pomegranate Council",
        "https://pomegranates.org/about/why-california/",
        "",
        "High",
        "",
        "",
    ),
    (
        "B1",
        74,
        "Dưa lưới Nhật",
        "Trái cây",
        "Nhật Bản",
        "Prefecture",
        "Ibaraki",
        "Ibaraki melon shipping window",
        4,
        10,
        "5-9",
        "Yes",
        "Ibaraki melons are shipped from late April to late October; Andes, Quincy and Takami peak from spring to early summer while Earl's melons run from summer to autumn.",
        "Melons",
        "Ibaraki Prefecture",
        "https://exports.pref.ibaraki.jp/en/product/detail?id=714713",
        "",
        "High",
        "",
        "",
    ),
    (
        "B1",
        74,
        "Dưa lưới Nhật",
        "Trái cây",
        "Nhật Bản",
        "Prefecture",
        "Shizuoka",
        "Greenhouse aroma melon supply",
        1,
        12,
        "Year-round",
        "No",
        "Shizuoka's premium aroma melon is cultivated in greenhouses and described as available throughout the year.",
        "Aroma Melon",
        "Shizuoka Prefecture Food Information Center",
        "https://fujinokuni.shokunomiyako-shizuoka.pref.shizuoka.jp/en/buy-food/1884",
        "",
        "High",
        "Catalog origin is generic Japan, so exact supplied prefecture still needs transactional confirmation.",
        "Good example of locality-specific difference versus open-field/standard shipping calendars.",
    ),
    (
        "B1",
        75,
        "Mận đỏ Mỹ",
        "Trái cây",
        "Mỹ",
        "State",
        "California",
        "California plum season",
        5,
        12,
        "6-8",
        "Yes",
        "The California plum fact sheet describes the growing season as May-December with the heaviest volume June-August.",
        "Plums",
        "California Department of Education",
        "https://www.cde.ca.gov/ls/nu/fd/plum.asp",
        "",
        "High",
        "",
        "",
    ),
    (
        "B1",
        76,
        "Lê Nam Phi",
        "Trái cây",
        "Nam Phi",
        "Province / region",
        "Western Cape",
        "South African pear harvest start",
        1,
        "",
        "",
        "No",
        "Hortgro notes that early pear cultivars are harvested in January; exact end windows vary by cultivar.",
        "Pome Fruit Industry Estimates 2025 Season",
        "Hortgro",
        "https://www.hortgro.co.za/wp-content/uploads/docs/dlm_uploads/2025/01/Pome-fruit-industry-estimates_2025-SEASON_FINAL.pdf",
        "2025-01",
        "High",
        "Generic South African pear timing still needs cultivar-specific end-month confirmation.",
        "Use January as verified harvest start only.",
    ),
    (
        "B1",
        77,
        "Kiwi vàng New Zealand",
        "Trái cây",
        "New Zealand",
        "Country",
        "New Zealand",
        "New Zealand kiwifruit harvest",
        3,
        5,
        "3-5",
        "Yes",
        "Zespri states New Zealand kiwifruit is harvested from March through May.",
        "When is Kiwifruit Season?",
        "Zespri",
        "https://www.zespri.com/en-US/blogdetail/when-is-kiwifruit-season",
        "",
        "High",
        "",
        "",
    ),
    (
        "B1",
        78,
        "Dứa MD2 Thái Lan",
        "Trái cây",
        "Thái Lan",
        "Country",
        "Thailand",
        "Thailand pineapple general availability",
        1,
        12,
        "4-7",
        "No",
        "Thailand DOAE lists pineapple as available year-round, with peak volume in April-July.",
        "Pineapple",
        "Department of Agricultural Extension, Thailand",
        "https://www.doae.go.th/en/pineapple-2/",
        "",
        "High",
        "The official source is for Thai pineapple generally, not a month-by-month MD2-specific calendar.",
        "Keep the MD2 varietal season flagged as provisional.",
    ),
]

RAW_RESEARCH_ROWS = [
    ("B1", 61, "Táo Fuji Nhật", "Origin country", "Japan", "Product origin", "Fruit and Vegetables: Apples", "Japan Food Export Fair / Aomori Apple Council", "https://j-fec.com/en/producer/pdf/en/20/", "2022", "High", "", ""),
    ("B1", 61, "Táo Fuji Nhật", "Origin province/region", "Aomori Prefecture", "Product origin", "Fruit and Vegetables: Apples", "Japan Food Export Fair / Aomori Apple Council", "https://j-fec.com/en/producer/pdf/en/20/", "2022", "High", "", ""),
    ("B1", 61, "Táo Fuji Nhật", "Flavor / texture", "Juicy flesh; fragrant, sweet profile; vivid red skin", "Commodity description", "Fruit and Vegetables: Apples", "Japan Food Export Fair / Aomori Apple Council", "https://j-fec.com/en/producer/pdf/en/20/", "2022", "High", "", ""),
    ("B1", 61, "Táo Fuji Nhật", "Preservation guidance", "Strict temperature control and storage allow quality to remain strong beyond New Year.", "Commodity handling", "Fruit and Vegetables: Apples", "Japan Food Export Fair / Aomori Apple Council", "https://j-fec.com/en/producer/pdf/en/20/", "2022", "High", "", ""),
    ("B1", 62, "Cam Cara ruột đỏ", "Origin country", "Australia", "Product origin", "What's in Season", "Citrus Australia", "https://citrusaustralia.com.au/buying-australian/season/", "", "High", "", ""),
    ("B1", 62, "Cam Cara ruột đỏ", "Origin regions", "Murray Valley, Riverina NSW, Riverland SA", "Australian citrus regions", "What's in Season", "Citrus Australia", "https://citrusaustralia.com.au/buying-australian/season/", "", "High", "", ""),
    ("B1", 62, "Cam Cara ruột đỏ", "Variety note", "Cara Cara is an early-to-mid-season navel orange with sweet, red flesh and low/absent seeds.", "Variety description", "Citrus Traceability Pilot", "Agriculture Victoria", "https://agriculture.vic.gov.au/__data/assets/pdf_file/0010/629902/Citrus-Traceability-Pilot.pdf", "", "High", "", ""),
    ("B1", 63, "Nho Mỹ không hạt", "Origin country", "United States", "Product origin", "Grapes From California Merchandising and Training Guide", "California Table Grape Commission", "https://www.grapesfromcalifornia.com/wp-content/uploads/2023/06/20230605-grapes-from-california-merchandising-and-training-guide-2023.pdf", "2023-06-05", "High", "", ""),
    ("B1", 63, "Nho Mỹ không hạt", "Origin province/region", "California - Coachella Valley, San Joaquin Valley", "Commodity production areas", "Grapes From California Merchandising and Training Guide", "California Table Grape Commission", "https://www.grapesfromcalifornia.com/wp-content/uploads/2023/06/20230605-grapes-from-california-merchandising-and-training-guide-2023.pdf", "2023-06-05", "High", "", ""),
    ("B1", 71, "Măng cụt", "Origin country", "Thailand", "Product origin", "Mangosteen", "Department of Agricultural Extension, Thailand", "https://www.doae.go.th/en/mangosteen-2/", "", "High", "", ""),
    ("B1", 71, "Măng cụt", "Flavor / texture", "Sweet and juicy flesh with mild fragrance.", "Commodity description", "Mangosteen", "Department of Agricultural Extension, Thailand", "https://www.doae.go.th/en/mangosteen-2/", "", "High", "", ""),
    ("B1", 71, "Măng cụt", "Preservation guidance", "3-4 weeks at 13°C and 90-95% RH.", "Commodity handling", "Mangosteen", "Department of Agricultural Extension, Thailand", "https://www.doae.go.th/en/mangosteen-2/", "", "High", "", ""),
    ("B1", 73, "Lựu đỏ Mỹ", "Origin country", "United States", "Product origin", "Why California", "POM Wonderful / California Pomegranate Council", "https://pomegranates.org/about/why-california/", "", "High", "", ""),
    ("B1", 73, "Lựu đỏ Mỹ", "Origin province/region", "California - San Joaquin Valley", "Commodity production area", "Why California", "POM Wonderful / California Pomegranate Council", "https://pomegranates.org/about/why-california/", "", "High", "", ""),
    ("B1", 74, "Dưa lưới Nhật", "Origin country", "Japan", "Product origin", "Melons", "Ibaraki Prefecture", "https://exports.pref.ibaraki.jp/en/product/detail?id=714713", "", "High", "", ""),
    ("B1", 74, "Dưa lưới Nhật", "Origin province/region", "Ibaraki / Shizuoka", "Verified localities from official sources", "Aroma Melon", "Shizuoka Prefecture Food Information Center", "https://fujinokuni.shokunomiyako-shizuoka.pref.shizuoka.jp/en/buy-food/1884", "", "High", "Catalog origin remains generic Japan.", ""),
    ("B1", 74, "Dưa lưới Nhật", "Flavor / texture", "Rich aroma, sweetness and soft texture.", "Commodity description", "Aroma Melon", "Shizuoka Prefecture Food Information Center", "https://fujinokuni.shokunomiyako-shizuoka.pref.shizuoka.jp/en/buy-food/1884", "", "High", "", ""),
    ("B1", 74, "Dưa lưới Nhật", "Preservation guidance", "Keep refrigerated.", "Commodity handling", "Melons", "Ibaraki Prefecture", "https://exports.pref.ibaraki.jp/en/product/detail?id=714713", "", "High", "", ""),
    ("B1", 75, "Mận đỏ Mỹ", "Origin country", "United States", "Product origin", "Plums", "California Department of Education", "https://www.cde.ca.gov/ls/nu/fd/plum.asp", "", "High", "", ""),
    ("B1", 75, "Mận đỏ Mỹ", "Origin province/region", "California", "Commodity production area", "Plums", "California Department of Education", "https://www.cde.ca.gov/ls/nu/fd/plum.asp", "", "High", "", ""),
    ("B1", 76, "Lê Nam Phi", "Origin country", "South Africa", "Product origin", "Pome Fruit Industry Estimates 2025 Season", "Hortgro", "https://www.hortgro.co.za/wp-content/uploads/docs/dlm_uploads/2025/01/Pome-fruit-industry-estimates_2025-SEASON_FINAL.pdf", "2025-01", "High", "", ""),
    ("B1", 76, "Lê Nam Phi", "Origin province/region", "Western Cape - Ceres, Groenland, Wolseley/Tulbagh", "Main pear producing areas", "Pear Market Value Chain Profile 2020", "Department of Agriculture, Land Reform and Rural Development (South Africa)", "https://www.nda.gov.za/images/Branches/Economica%20Development%20Trade%20and%20Marketing/marketing/annual-publications/fruits/pear-market-value-chain-profile-2020.pdf", "2020", "High", "", ""),
    ("B1", 77, "Kiwi vàng New Zealand", "Origin country", "New Zealand", "Product origin", "When is Kiwifruit Season?", "Zespri", "https://www.zespri.com/en-US/blogdetail/when-is-kiwifruit-season", "", "High", "", ""),
    ("B1", 78, "Dứa MD2 Thái Lan", "Origin country", "Thailand", "Product origin", "Pineapple", "Department of Agricultural Extension, Thailand", "https://www.doae.go.th/en/pineapple-2/", "", "High", "", ""),
    ("B1", 78, "Dứa MD2 Thái Lan", "Variety note", "MD2 is one of Thailand's main fresh-export pineapple varieties.", "Varietal context", "An Economic Study of Fresh Pineapple Production of Pineapple cv. MD2 in Phetchaburi Province", "Department of Agriculture, Thailand", "https://info.doa.go.th/research/index.php?act=view_detail&id=4460", "2020", "Medium", "Useful to confirm the MD2 variety context, but not a direct national season calendar.", ""),
]


def insert_value_at(row, index, value):
    values = list(row)
    values.insert(index, value)
    return tuple(values)


SEASONALITY_ROWS = [insert_value_at(row, 18, BATCH1_COLLECTION_TIMESTAMP) for row in RAW_SEASONALITY_ROWS]
RESEARCH_ROWS = [insert_value_at(row, 11, BATCH1_COLLECTION_TIMESTAMP) for row in RAW_RESEARCH_ROWS]

PRODUCT_SUMMARIES = {
    61: {
        "Review Status": "Partially verified - batch 1",
        "Verified Seasonality Summary": "Aomori (Nhật): Fuji nổi bật từ khoảng tháng 11; xuất khẩu táo Aomori bắt đầu từ tháng 10 và bảo quản lạnh giúp kéo dài chất lượng qua đầu năm.",
        "Verified Origin Country": "Japan",
        "Verified Origin Province/Region": "Aomori Prefecture",
        "Verified Domestic/Imported": "Imported",
        "Verified Preservation Guidance": "Strict temperature control; quality can extend beyond New Year in storage.",
        "Verified Flavor Profile": "Juicy, fragrant, sweet",
        "Verified Texture Profile": "Crisp / juicy",
        "Primary Source URL": "https://j-fec.com/en/producer/pdf/en/20/",
        "Secondary Source URL": "",
        "Thời gian thu thập": BATCH1_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "Tháng kết thúc đang suy từ note bảo quản qua đầu năm; chưa có closing month Fuji thật sự rõ trong nguồn chính.",
    },
    62: {
        "Review Status": "Partially verified - batch 1",
        "Verified Seasonality Summary": "Australia: nhóm navel vào mùa tháng 6-10; Cara Cara được mô tả là giống early-to-mid-season trong nhóm navel.",
        "Verified Origin Country": "Australia",
        "Verified Origin Province/Region": "Murray Valley / Riverina / Riverland",
        "Verified Domestic/Imported": "Imported",
        "Verified Flavor Profile": "Sweet, red-fleshed",
        "Verified Texture Profile": "Juicy / low-seed",
        "Primary Source URL": "https://citrusaustralia.com.au/buying-australian/season/",
        "Secondary Source URL": "https://agriculture.vic.gov.au/__data/assets/pdf_file/0010/629902/Citrus-Traceability-Pilot.pdf",
        "Thời gian thu thập": BATCH1_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "Khung tháng Cara Cara đang suy trong cửa sổ navel June-October vì chưa có lịch Cara Cara Australia theo tháng từ nguồn chính thống mạnh hơn.",
    },
    63: {
        "Review Status": "Partially verified - batch 1",
        "Verified Seasonality Summary": "California, Mỹ: Coachella Valley từ tháng 5 đến đầu tháng 7; San Joaquin Valley từ đầu tháng 7 đến tháng 12; nho California có mặt trên thị trường từ tháng 5 đến tháng 1.",
        "Verified Origin Country": "United States",
        "Verified Origin Province/Region": "California - Coachella Valley / San Joaquin Valley",
        "Verified Domestic/Imported": "Imported",
        "Primary Source URL": "https://www.grapesfromcalifornia.com/wp-content/uploads/2023/06/20230605-grapes-from-california-merchandising-and-training-guide-2023.pdf",
        "Secondary Source URL": "",
        "Thời gian thu thập": BATCH1_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "",
    },
    71: {
        "Review Status": "Partially verified - batch 1",
        "Verified Seasonality Summary": "Thái Lan: sẵn hàng từ tháng 4-12, đỉnh tháng 4-8; tại Khiriwong (Nakhon Si Thammarat) mùa trái tập trung khoảng tháng 7-9.",
        "Verified Origin Country": "Thailand",
        "Verified Domestic/Imported": "Imported",
        "Verified Preservation Guidance": "3-4 weeks at 13°C and 90-95% RH.",
        "Verified Flavor Profile": "Sweet, juicy, mild fragrance",
        "Verified Texture Profile": "Soft segmented flesh",
        "Primary Source URL": "https://www.doae.go.th/en/mangosteen-2/",
        "Secondary Source URL": "https://www.thailand.go.th/issue-focus-detail/001_02_255",
        "Thời gian thu thập": BATCH1_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "",
    },
    73: {
        "Review Status": "Partially verified - batch 1",
        "Verified Seasonality Summary": "California (San Joaquin Valley): mùa lựu từ khoảng tháng 10 đến tháng 1.",
        "Verified Origin Country": "United States",
        "Verified Origin Province/Region": "California - San Joaquin Valley",
        "Verified Domestic/Imported": "Imported",
        "Primary Source URL": "https://pomegranates.org/about/why-california/",
        "Secondary Source URL": "",
        "Thời gian thu thập": BATCH1_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "",
    },
    74: {
        "Review Status": "Partially verified - batch 1",
        "Verified Seasonality Summary": "Nhật Bản: Ibaraki shipping từ cuối tháng 4 đến cuối tháng 10; Shizuoka greenhouse có thể cung ứng quanh năm.",
        "Verified Origin Country": "Japan",
        "Verified Origin Province/Region": "Ibaraki / Shizuoka",
        "Verified Domestic/Imported": "Imported",
        "Verified Preservation Guidance": "Keep refrigerated.",
        "Verified Flavor Profile": "Rich aroma, sweet",
        "Verified Texture Profile": "Soft / juicy flesh",
        "Primary Source URL": "https://exports.pref.ibaraki.jp/en/product/detail?id=714713",
        "Secondary Source URL": "https://fujinokuni.shokunomiyako-shizuoka.pref.shizuoka.jp/en/buy-food/1884",
        "Thời gian thu thập": BATCH1_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "Catalog chỉ ghi Japan; để gắn tag mùa vụ thật chính xác cần biết lô hàng thực đến từ Ibaraki hay Shizuoka/nhà kính khác.",
    },
    75: {
        "Review Status": "Partially verified - batch 1",
        "Verified Seasonality Summary": "California: mùa mận kéo dài từ tháng 5 đến tháng 12, sản lượng mạnh nhất vào tháng 6-8.",
        "Verified Origin Country": "United States",
        "Verified Origin Province/Region": "California",
        "Verified Domestic/Imported": "Imported",
        "Primary Source URL": "https://www.cde.ca.gov/ls/nu/fd/plum.asp",
        "Secondary Source URL": "",
        "Thời gian thu thập": BATCH1_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "",
    },
    76: {
        "Review Status": "Partially verified - batch 1",
        "Verified Seasonality Summary": "Nam Phi: mùa lê được xác nhận bắt đầu thu hoạch từ tháng 1; vùng trồng chính nằm ở Western Cape như Ceres, Groenland và Wolseley/Tulbagh.",
        "Verified Origin Country": "South Africa",
        "Verified Origin Province/Region": "Western Cape - Ceres / Groenland / Wolseley-Tulbagh",
        "Verified Domestic/Imported": "Imported",
        "Primary Source URL": "https://www.hortgro.co.za/wp-content/uploads/docs/dlm_uploads/2025/01/Pome-fruit-industry-estimates_2025-SEASON_FINAL.pdf",
        "Secondary Source URL": "https://www.nda.gov.za/images/Branches/Economica%20Development%20Trade%20and%20Marketing/marketing/annual-publications/fruits/pear-market-value-chain-profile-2020.pdf",
        "Thời gian thu thập": BATCH1_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "Nguồn mạnh mới xác nhận tháng bắt đầu và vùng trồng chính; khung kết thúc của generic South African pear vẫn cần kiểm tra thêm theo cultivar.",
    },
    77: {
        "Review Status": "Partially verified - batch 1",
        "Verified Seasonality Summary": "New Zealand: thu hoạch kiwi từ tháng 3 đến tháng 5.",
        "Verified Origin Country": "New Zealand",
        "Verified Domestic/Imported": "Imported",
        "Primary Source URL": "https://www.zespri.com/en-US/blogdetail/when-is-kiwifruit-season",
        "Secondary Source URL": "",
        "Thời gian thu thập": BATCH1_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "",
    },
    78: {
        "Review Status": "Partially verified - batch 1",
        "Verified Seasonality Summary": "Thái Lan: dứa nói chung có quanh năm, đỉnh tháng 4-7; lịch riêng cho MD2 vẫn cần nguồn chính thống mạnh hơn.",
        "Verified Origin Country": "Thailand",
        "Verified Domestic/Imported": "Imported",
        "Primary Source URL": "https://www.doae.go.th/en/pineapple-2/",
        "Secondary Source URL": "https://info.doa.go.th/research/index.php?act=view_detail&id=4460",
        "Thời gian thu thập": BATCH1_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "Mùa vụ hiện mới xác nhận ở mức pineapple Thailand general; chưa có lịch tháng đủ mạnh riêng cho MD2 Thailand.",
    },
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--workbook",
        default=str(DEFAULT_WORKBOOK_PATH),
        help="Workbook path to update.",
    )
    return parser.parse_args()


def normalize_row(row):
    return tuple(None if cell == "" else cell for cell in row)


def dedupe_sheet(ws):
    seen = set()
    deduped = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        normalized = normalize_row(row)
        if not any(cell is not None for cell in normalized):
            continue
        if normalized in seen:
            continue
        seen.add(normalized)
        deduped.append(normalized)

    ws.delete_rows(2, ws.max_row)
    for row in deduped:
        ws.append(row)


def append_unique_rows(ws, rows):
    existing = {
        normalize_row(row)
        for row in ws.iter_rows(min_row=2, values_only=True)
        if any(cell is not None and cell != "" for cell in row)
    }
    for row in rows:
        normalized = normalize_row(row)
        if normalized not in existing:
            ws.append(normalized)
            existing.add(normalized)


def update_products_review(ws):
    headers = [cell.value for cell in ws[1]]
    header_index = {header: idx + 1 for idx, header in enumerate(headers)}

    product_row_map = {}
    for row in range(2, ws.max_row + 1):
        product_id = ws.cell(row=row, column=1).value
        if product_id is not None:
            product_row_map[product_id] = row

    for product_id, updates in PRODUCT_SUMMARIES.items():
        row = product_row_map.get(product_id)
        if row is None:
            continue
        for header, value in updates.items():
            col = header_index[header]
            ws.cell(row=row, column=col).value = value


def main() -> None:
    args = parse_args()
    workbook_path = Path(args.workbook)
    wb = load_workbook(workbook_path)

    dedupe_sheet(wb["Seasonality Evidence"])
    dedupe_sheet(wb["Research Evidence"])
    append_unique_rows(wb["Seasonality Evidence"], SEASONALITY_ROWS)
    append_unique_rows(wb["Research Evidence"], RESEARCH_ROWS)
    update_products_review(wb["Products Review"])

    wb.save(workbook_path)


if __name__ == "__main__":
    main()

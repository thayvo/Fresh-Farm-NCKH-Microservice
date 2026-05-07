from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK_PATH = REPO_ROOT / "output" / "spreadsheet" / "recommendation_product_data_review.xlsx"
ROUND4_COLLECTION_TIMESTAMP = "2026-04-24 21:05:08 ICT"


SEASONALITY_ROWS = [
    (
        "B3",
        69,
        "Đu đủ vàng",
        "Trái cây",
        "Bình Phước, Việt Nam",
        "Province",
        "Bình Phước",
        "Cây ăn trái nhiệt đới gần như quanh năm",
        1,
        12,
        "Year-round",
        "No",
        "Cổng thông tin Bình Phước ghi nhiều chủng loại cây ăn trái của tỉnh hầu như cho thu hoạch quanh năm, trong danh sách có đu đủ.",
        "Bình Phước có nhiều chủng loại cây ăn trái hầu như cho thu hoạch quanh năm",
        "Cổng thông tin điện tử tỉnh Bình Phước",
        "https://binhphuoc.gov.vn/vi/news/Hoat-dong-huyen-thi/phuoc-co-nhieu-chung-loai-cay-an-trai-hau-nhu-cho-thu-hoach-quanh-nam-21193.html",
        "",
        "High",
        ROUND4_COLLECTION_TIMESTAMP,
        "Nguồn xác nhận đu đủ trong nhóm cây ăn trái Bình Phước thu hoạch gần như quanh năm, chưa tách riêng giống đu đủ vàng.",
        "",
    ),
]

RESEARCH_ROWS = [
    (
        "B3",
        69,
        "Đu đủ vàng",
        "Origin/seasonality context",
        "Bình Phước có đu đủ trong nhóm cây ăn trái hầu như cho thu hoạch quanh năm.",
        "Bình Phước",
        "Bình Phước có nhiều chủng loại cây ăn trái hầu như cho thu hoạch quanh năm",
        "Cổng thông tin điện tử tỉnh Bình Phước",
        "https://binhphuoc.gov.vn/vi/news/Hoat-dong-huyen-thi/phuoc-co-nhieu-chung-loai-cay-an-trai-hau-nhu-cho-thu-hoach-quanh-nam-21193.html",
        "",
        "High",
        ROUND4_COLLECTION_TIMESTAMP,
        "Chưa tìm được nguồn riêng cho giống đu đủ vàng tại Bình Phước.",
        "",
    ),
]

PRODUCT_SUMMARIES = {
    69: {
        "Review Status": "Partially verified - batch 3",
        "Verified Seasonality Summary": "Bình Phước: đu đủ nằm trong nhóm cây ăn trái hầu như cho thu hoạch quanh năm; chưa tách riêng giống đu đủ vàng.",
        "Verified Origin Country": "Vietnam",
        "Verified Origin Province/Region": "Bình Phước",
        "Verified Domestic/Imported": "Domestic",
        "Primary Source URL": "https://binhphuoc.gov.vn/vi/news/Hoat-dong-huyen-thi/phuoc-co-nhieu-chung-loai-cay-an-trai-hau-nhu-cho-thu-hoach-quanh-nam-21193.html",
        "Secondary Source URL": "",
        "Thời gian thu thập": ROUND4_COLLECTION_TIMESTAMP,
        "Tôi chưa chắc chắn": "Nguồn nói đu đủ trong nhóm cây ăn trái Bình Phước nói chung, chưa có lịch riêng cho giống đu đủ vàng.",
    },
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK_PATH))
    return parser.parse_args()


def normalize_row(row):
    return tuple(None if cell == "" else cell for cell in row)


def append_unique_rows(ws, rows):
    existing = {
        normalize_row(row)
        for row in ws.iter_rows(min_row=2, values_only=True)
        if any(cell is not None and cell != "" for cell in row)
    }
    for row in rows:
        normalized = normalize_row(row)
        if normalized not in existing:
            ws.append(normalized)
            existing.add(normalized)


def update_products_review(ws):
    headers = [cell.value for cell in ws[1]]
    header_index = {header: idx + 1 for idx, header in enumerate(headers)}
    product_row_map = {}
    for row in range(2, ws.max_row + 1):
        product_id = ws.cell(row=row, column=1).value
        if product_id is not None:
            product_row_map[product_id] = row

    for product_id, updates in PRODUCT_SUMMARIES.items():
        row = product_row_map.get(product_id)
        if row is None:
            continue
        for header, value in updates.items():
            ws.cell(row=row, column=header_index[header]).value = value


def main() -> None:
    args = parse_args()
    wb = load_workbook(Path(args.workbook))
    append_unique_rows(wb["Seasonality Evidence"], SEASONALITY_ROWS)
    append_unique_rows(wb["Research Evidence"], RESEARCH_ROWS)
    update_products_review(wb["Products Review"])
    wb.save(Path(args.workbook))


if __name__ == "__main__":
    main()

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


REPO_ROOT = Path(__file__).resolve().parents[2]
SOURCE_WORKBOOK = REPO_ROOT / "output" / "spreadsheet" / "catalog_product_origin_single_sheet.xlsx"
OUTPUT_WORKBOOK = REPO_ROOT / "output" / "spreadsheet" / "catalog_product_seasonality_single_sheet.xlsx"

UPDATED_AT = "2026-04-24 21:45:00 ICT"


HEADERS = [
    "SeasonRowID",
    "ProductID",
    "Quốc gia",
    "Tỉnh/Thành/Vùng",
    "Khu vực chi tiết",
    "SeasonType",
    "SeasonLabel",
    "StartMonth",
    "EndMonth",
    "PeakMonths",
    "IsYearRound",
    "HasPeakSeason",
    "IsControlledCultivation",
    "IsImportedSeason",
    "IsOffSeason",
    "IsPostHarvestAvailability",
    "SeasonScoreWeight",
    "ConfidenceLevel",
    "Nguồn/Ghi chú",
    "Thời gian cập nhật",
]


# Product-specific season windows. Each tuple:
# season_type, label, start_month, end_month, peak_months, is_year_round, db_hint, note
PRODUCT_SEASONS = {
    46: [
        ("Vụ chính", "Khoai lang Nhật vụ Xuân - Bắc Giang", 2, 5, "4-5", "No", "Theo mùa", "Nguồn khuyến nông: trồng sau Tết, thu đầu tháng 5 âm lịch."),
        ("Vụ chính", "Khoai lang Nhật vụ Đông - Bắc Giang", 8, 12, "11-12", "No", "Theo mùa", "Nguồn khuyến nông: trồng cuối 8-đầu 9, thu tháng 12 âm lịch."),
        ("Vụ tham khảo", "Khoai lang tím Nhật vụ Đông Xuân - Tây Ninh", 12, 4, "2-4", "No", "Theo mùa qua năm", "Nguồn khuyến nông Tây Ninh; dùng làm tham khảo cho giống Nhật trồng VN."),
    ],
    61: [
        ("Vụ chính", "Táo Fuji Nhật vụ thu hoạch", 9, 12, "10-11", "No", "Theo mùa", "Aomori/Japan apple harvest season."),
        ("Sau thu hoạch", "Táo Fuji bảo quản/lưu kho", 1, 4, "1-3", "No", "Có sau thu hoạch", "Táo Fuji có khả năng bảo quản sau vụ thu hoạch."),
    ],
    62: [("Vụ nhập khẩu", "Cam Cara/Navel Úc", 6, 10, "7-9", "No", "Theo mùa", "Lịch citrus Australia tham khảo.")],
    63: [("Vụ nhập khẩu", "Nho California", 5, 1, "7-11", "No", "Theo mùa qua năm", "California table grapes thường có mùa từ cuối xuân đến đầu đông.")],
    64: [
        ("Vụ chính", "Xoài cát Hòa Lộc vụ thuận", 3, 5, "4-5", "No", "Theo mùa", "Mùa xoài chính miền Tây/Tiền Giang."),
        ("Vụ phụ/rải vụ", "Xoài cát Hòa Lộc rải vụ", 10, 12, "11-12", "No", "Vụ phụ", "Nguồn Bộ Công Thương có nhắc mùa tăng thêm cuối năm."),
    ],
    65: [
        ("Vụ chính", "Dưa hấu Long An vụ Tết", 12, 2, "1-2", "No", "Theo mùa qua năm", "Dưa hấu Long An phục vụ thị trường Tết."),
        ("Vụ phụ", "Dưa hấu Long An vụ hè", 4, 7, "5-6", "No", "Theo mùa", "Vụ hè tham khảo cho dưa hấu miền Nam."),
    ],
    66: [("Vụ chính", "Bơ Đắk Lắk/Booth", 6, 11, "8-10", "No", "Theo mùa", "Bơ Đắk Lắk thường tập trung giữa năm đến cuối năm; Booth muộn hơn một số giống.")],
    67: [
        ("Vụ chính", "Dâu tây Đà Lạt mùa mát", 11, 4, "12-3", "No", "Theo mùa qua năm", "Dâu tây Đà Lạt thuận mùa mát."),
        ("Vụ phụ/nhà kính", "Dâu tây Đà Lạt rải vụ", 5, 10, "6-8", "No", "Rải vụ có kiểm soát", "Có thể rải vụ/nhà kính tùy giống và trang trại."),
    ],
    68: [("Quanh năm", "Chuối già Nam Mỹ", 1, 12, "Quanh năm", "Yes", "Quanh năm", "Chuối trồng vùng nhiệt đới có thể thu quanh năm theo lứa.")],
    69: [("Quanh năm", "Đu đủ Bình Phước", 1, 12, "Quanh năm", "Yes", "Quanh năm", "Cổng thông tin Bình Phước ghi nhóm cây ăn trái có đu đủ thu gần quanh năm.")],
    70: [
        ("Vụ thuận", "Thanh long Bình Thuận vụ thuận", 4, 9, "5-8", "No", "Theo mùa", "Thanh long Bình Thuận vụ thuận theo điều kiện tự nhiên."),
        ("Vụ nghịch", "Thanh long Bình Thuận chong đèn/vụ nghịch", 10, 3, "11-2", "No", "Vụ nghịch qua năm", "Bình Thuận có sản xuất vụ nghịch bằng xử lý ra hoa/chong đèn."),
    ],
    71: [
        ("Vụ chính", "Măng cụt Thái Lan miền Đông", 4, 7, "5-6", "No", "Theo mùa", "Chanthaburi/Rayong/Trat thường là vùng măng cụt chính."),
        ("Vụ phụ/miền Nam", "Măng cụt Thái Lan kéo dài/trái vụ", 8, 10, "8-9", "No", "Vụ phụ", "Một số vùng miền Nam Thái Lan có thể kéo dài mùa."),
    ],
    72: [("Quanh năm/rải vụ", "Mít Thái Tiền Giang", 1, 12, "3-7; 10-12", "Yes", "Quanh năm, có mùa rộ", "Mít Thái có thể rải vụ, vùng Tiền Giang thu nhiều đợt trong năm.")],
    73: [("Vụ nhập khẩu", "Lựu đỏ California", 8, 11, "9-10", "No", "Theo mùa", "California pomegranate season.")],
    74: [
        ("Vụ chính", "Dưa lưới Nhật Ibaraki/Shizuoka", 5, 8, "6-7", "No", "Theo mùa", "Dưa lưới Nhật có mùa mạnh cuối xuân-hè, tùy vùng/nhà kính."),
        ("Nhà kính", "Dưa lưới Nhật nhà kính", 9, 12, "10-11", "No", "Cung ứng có kiểm soát", "Nhà kính có thể mở rộng mùa cung ứng."),
    ],
    75: [("Vụ nhập khẩu", "Mận đỏ California", 5, 10, "6-8", "No", "Theo mùa", "California plum season.")],
    76: [("Vụ nhập khẩu", "Lê Nam Phi", 2, 8, "3-6", "No", "Theo mùa", "South Africa pear export/harvest season tham khảo.")],
    77: [("Vụ nhập khẩu", "Kiwi vàng New Zealand", 4, 11, "5-9", "No", "Theo mùa", "New Zealand kiwifruit harvest/export season, Bay of Plenty là vùng chính.")],
    78: [("Quanh năm", "Dứa MD2 Thái Lan", 1, 12, "3-6; 10-12", "Yes", "Quanh năm, có mùa rộ", "Dứa Thái có cung ứng quanh năm, Prachuap Khiri Khan là vùng tham khảo.")],
    116: [("Quanh năm", "Chuối già Việt Nam Tiền Giang", 1, 12, "Quanh năm", "Yes", "Quanh năm", "Chuối có thể thu quanh năm theo lứa ở vùng nhiệt đới.")],
    117: [("Quanh năm", "Ổi xá lị Long An", 1, 12, "3-6; 9-12", "Yes", "Quanh năm, có mùa rộ", "Ổi có thể xử lý/rải vụ, thường có nhiều đợt trong năm.")],
    118: [
        ("Vụ chính", "Dưa hấu mini Long An vụ Tết", 12, 2, "1-2", "No", "Theo mùa qua năm", "Dưa hấu Long An phục vụ Tết."),
        ("Vụ phụ", "Dưa hấu mini Long An vụ hè", 4, 7, "5-6", "No", "Theo mùa", "Vụ hè tham khảo."),
    ],
    119: [
        ("Vụ thuận", "Thanh long ruột trắng Bình Thuận vụ thuận", 4, 9, "5-8", "No", "Theo mùa", "Thanh long Bình Thuận vụ thuận."),
        ("Vụ nghịch", "Thanh long ruột trắng Bình Thuận vụ nghịch", 10, 3, "11-2", "No", "Vụ nghịch qua năm", "Sản xuất vụ nghịch bằng chong đèn/xử lý ra hoa."),
    ],
    120: [("Vụ chính", "Sầu riêng Đắk Lắk/Ri6", 7, 10, "8-9", "No", "Theo mùa", "Sầu riêng Đắk Lắk thường thu tập trung giữa-cuối năm.")],
}


COOL_SEASON_KEYWORDS = {
    "Cải bó xôi",
    "Cải thìa",
    "Bông cải xanh",
    "Súp lơ trắng",
    "Rau cải ngọt",
    "Cải ngồng",
    "Cải xoăn kale",
    "Bông cải trắng",
    "Bắp cải tím",
    "Cải Brussels sprouts",
    "Xà lách xoăn",
    "Xà lách lô lô",
    "Cải bẹ xanh",
    "Bông cải xanh baby",
    "Bông cải trắng mini",
    "Su hào",
    "Củ cải trắng",
    "Củ cải đỏ",
    "Cà rốt",
    "Khoai tây vàng",
    "Củ dền",
}


WARM_VEG_KEYWORDS = {
    "Rau muống",
    "Rau lang",
    "Rau dền đỏ",
    "Rau diếp cá",
    "Rau má",
    "Rau mồng tơi",
    "Rau húng quế",
    "Rau húng",
    "Húng lủi",
    "Tía tô",
    "Rau răm",
    "Ngò rí",
    "Ngò gai",
    "Hành lá",
    "Ớt hiểm",
    "Bầu xanh",
    "Mướp hương",
    "Đậu bắp",
    "Bí đỏ",
    "Bí đao",
    "Đậu que",
    "Đậu cove",
    "Cà tím",
}


CONTROLLED_OR_SHORT_CYCLE = {
    "Rau mầm hướng dương",
    "Rau mầm đậu Hà Lan",
    "Rau mầm alfalfa",
    "Nấm đùi gà",
    "Nấm mỡ",
    "Nấm kim châm",
    "Nấm bào ngư",
    "Nấm rơm tươi",
    "Nấm đùi gà tươi",
    "Nấm hương tươi",
    "Nấm hương khô",
    "Nấm linh chi đỏ",
}


SPECIAL_VEG_SEASONS = {
    "Măng tây xanh": [("Vụ chính", "Măng tây Ninh Thuận mùa nắng", 2, 8, "3-7", "No", "Theo mùa", "Măng tây vùng khô nóng có thể thu nhiều lứa, thuận mùa nắng.")],
    "Hoa thiên lý": [("Vụ chính", "Hoa thiên lý mùa nóng/mưa", 4, 10, "5-8", "No", "Theo mùa", "Hoa thiên lý thường rộ khi thời tiết ấm.")],
    "Măng tre tươi": [("Vụ chính", "Măng tre mùa mưa", 5, 10, "6-8", "No", "Theo mùa", "Măng tre thường rộ mùa mưa.")],
    "Bông atiso": [("Vụ chính", "Atiso Đà Lạt mùa mát", 11, 4, "12-3", "No", "Theo mùa qua năm", "Atiso Đà Lạt thuận mùa mát.")],
    "Khoai lang tím": [("Vụ chính", "Khoai lang tím Đông Xuân", 12, 4, "2-4", "No", "Theo mùa qua năm", "Khoai lang thường có vụ Đông Xuân ở miền Nam.")],
    "Sắn lùn": [("Vụ chính", "Sắn Tây Nguyên", 11, 4, "12-3", "No", "Theo mùa qua năm", "Sắn thường thu sau mùa mưa/vào mùa khô.")],
    "Gừng tươi": [("Vụ chính", "Gừng Lâm Đồng", 10, 2, "11-1", "No", "Theo mùa qua năm", "Gừng thường thu khi cây già/cuối năm đến đầu năm.")],
    "Hành tây tím": [("Vụ chính", "Hành tây Đà Lạt mùa khô/mát", 12, 4, "1-3", "No", "Theo mùa qua năm", "Hành tây Đà Lạt thường thuận mùa mát.")],
    "Hành tỏi phi": [("Quanh năm", "Hành tỏi phi TP.HCM - hàng chế biến", 1, 12, "Quanh năm", "Yes", "Hàng chế biến quanh năm", "Hàng chế biến, không nên gán mùa vụ nông sản tươi.")],
    "Cà chua bi": [("Quanh năm", "Cà chua bi Đà Lạt", 1, 12, "11-4", "Yes", "Quanh năm, có mùa rộ", "Cà chua Đà Lạt có thể sản xuất quanh năm, thuận mùa mát.")],
    "Cà chua beef": [("Quanh năm", "Cà chua beef Lâm Đồng", 1, 12, "11-4", "Yes", "Quanh năm, có mùa rộ", "Cà chua Lâm Đồng có thể sản xuất quanh năm, thuận mùa mát.")],
    "Dưa leo mini": [("Quanh năm", "Dưa leo Bến Tre", 1, 12, "11-4", "Yes", "Quanh năm, có mùa rộ", "Dưa leo miền Nam có thể trồng nhiều vụ/năm.")],
    "Dưa chuột baby": [("Quanh năm", "Dưa chuột baby Lâm Đồng", 1, 12, "11-4", "Yes", "Quanh năm, có mùa rộ", "Dưa chuột có thể trồng nhiều vụ/năm.")],
    "Ớt chuông": [("Quanh năm", "Ớt chuông Đà Lạt", 1, 12, "11-4", "Yes", "Quanh năm, có mùa rộ", "Ớt chuông nhà kính/Đà Lạt có thể cung ứng quanh năm.")],
    "Tỏi tây": [("Vụ tham khảo", "Tỏi tây Trung Quốc mùa mát", 10, 4, "11-3", "No", "Theo mùa qua năm", "Tỏi tây là rau mùa mát; vùng Trung Quốc/tỉnh nhập khẩu cần xác minh.")],
    "khoai mỡ": [("Vụ chính", "Khoai mỡ Vĩnh Long", 2, 8, "7-8", "No", "Theo mùa", "Nguồn khuyến nông: xuống giống sau Đông Xuân, thu sau khoảng 6 tháng.")],
}


def default_seasons(product_name: str, category: str, province: str):
    if product_name in SPECIAL_VEG_SEASONS:
        return SPECIAL_VEG_SEASONS[product_name]

    if product_name in CONTROLLED_OR_SHORT_CYCLE or category == "Nấm":
        return [("Quanh năm", f"{product_name} - nuôi trồng/kiểm soát", 1, 12, "Quanh năm", "Yes", "Quanh năm có kiểm soát", "Nấm/rau mầm phụ thuộc chu kỳ nuôi trồng, không phải mùa tự nhiên cố định.")]

    if product_name in COOL_SEASON_KEYWORDS or ("Lâm Đồng" in province and category in {"Rau lá", "Rau ăn hoa / thân / mầm", "Củ & rễ"}):
        return [
            ("Vụ chính", f"{product_name} mùa mát", 11, 4, "12-3", "No", "Theo mùa qua năm", "Rau/củ khí hậu mát vùng cao nguyên, thuận vụ Đông-Xuân."),
            ("Vụ phụ", f"{product_name} rải vụ", 5, 10, "6-8", "No", "Rải vụ có kiểm soát", "Có thể rải vụ tại Đà Lạt/Lâm Đồng hoặc vùng phù hợp."),
        ]

    if product_name in WARM_VEG_KEYWORDS or category in {"Rau lá", "Rau thơm & gia vị", "Rau ăn quả"}:
        return [("Quanh năm", f"{product_name} vùng nhiệt đới", 1, 12, "5-10", "Yes", "Quanh năm, có mùa rộ", "Rau nhiệt đới/rau ăn quả có thể trồng nhiều vụ trong năm, thuận khi đủ nước.")]

    if category == "Củ & rễ":
        return [("Vụ chính", f"{product_name} vụ Đông-Xuân", 11, 4, "12-3", "No", "Theo mùa qua năm", "Củ/rễ thường thuận mùa khô/mát tùy vùng.")]

    return [("Quanh năm", f"{product_name} - tham khảo", 1, 12, "Quanh năm", "Yes", "Quanh năm tham khảo", "Chưa có rule riêng, tạm để quanh năm để không thiếu dữ liệu.")]


def load_products():
    wb = load_workbook(SOURCE_WORKBOOK, read_only=True, data_only=True)
    ws = wb["Danh sách sản phẩm"]
    headers = [cell.value for cell in ws[1]]
    idx = {header: i for i, header in enumerate(headers)}
    products = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        products.append(
            {
                "ProductID": row[idx["ProductID"]],
                "Tên sản phẩm": row[idx["Tên sản phẩm"]],
                "Loại sản phẩm": row[idx["Loại sản phẩm"]],
                "Quốc gia": row[idx["Quốc gia"]],
                "Tỉnh/Thành/Vùng": row[idx["Tỉnh/Thành/Vùng"]],
                "Khu vực chi tiết": row[idx["Khu vực chi tiết"]],
            }
        )
    return products


def build_rows(products):
    rows = []
    season_row_id = 1
    for product in products:
        product_id = int(product["ProductID"])
        seasons = PRODUCT_SEASONS.get(product_id)
        if seasons is None:
            seasons = default_seasons(product["Tên sản phẩm"], product["Loại sản phẩm"], product["Tỉnh/Thành/Vùng"])

        for season in seasons:
            season_type, label, start, end, peak, year_round, db_hint, note = season
            flags = derive_rule_flags(season_type, db_hint, year_round, peak)
            score_weight = derive_score_weight(season_type, flags)
            rows.append(
                [
                    season_row_id,
                    product_id,
                    product["Quốc gia"],
                    product["Tỉnh/Thành/Vùng"],
                    product["Khu vực chi tiết"] or "",
                    season_type,
                    label,
                    start,
                    end,
                    peak,
                    year_round,
                    flags["HasPeakSeason"],
                    flags["IsControlledCultivation"],
                    flags["IsImportedSeason"],
                    flags["IsOffSeason"],
                    flags["IsPostHarvestAvailability"],
                    score_weight,
                    "",
                    note,
                    UPDATED_AT,
                ]
            )
            season_row_id += 1
    return rows


def derive_score_weight(season_type: str, flags: dict[str, str]) -> int:
    if flags["HasPeakSeason"] == "Yes":
        return 80
    if season_type in {"Vụ chính", "Vụ thuận", "Vụ nhập khẩu"}:
        return 100
    if season_type == "Quanh năm":
        return 50
    if flags["IsOffSeason"] == "Yes":
        return 35
    if flags["IsControlledCultivation"] == "Yes":
        return 45
    if flags["IsPostHarvestAvailability"] == "Yes":
        return 30
    return 40


def derive_rule_flags(season_type: str, db_hint: str, year_round: str, peak_months: str):
    text = f"{season_type} {db_hint}".lower()
    has_peak = year_round == "Yes" and peak_months not in {"Quanh năm", "", None}
    return {
        "HasPeakSeason": "Yes" if has_peak else "No",
        "IsControlledCultivation": "Yes" if "kiểm soát" in text or "nhà kính" in text else "No",
        "IsImportedSeason": "Yes" if "nhập khẩu" in text else "No",
        "IsOffSeason": "Yes" if "vụ phụ" in text or "vụ nghịch" in text or "rải vụ" in text else "No",
        "IsPostHarvestAvailability": "Yes" if "sau thu hoạch" in text else "No",
    }


def format_workbook(ws):
    widths = {
        "A": 12,
        "B": 12,
        "C": 16,
        "D": 32,
        "E": 34,
        "F": 18,
        "G": 38,
        "H": 12,
        "I": 12,
        "J": 18,
        "K": 14,
        "L": 16,
        "M": 22,
        "N": 18,
        "O": 16,
        "P": 24,
        "Q": 18,
        "R": 18,
        "S": 58,
        "T": 22,
    }
    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(bold=True, color="FFFFFF")
    border = Border(bottom=Side(style="thin", color="D1D5DB"))
    fill_year = PatternFill("solid", fgColor="ECFDF5")
    fill_season = PatternFill("solid", fgColor="EFF6FF")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    header_index = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        is_year_round = row[header_index["IsYearRound"] - 1].value == "Yes"
        fill = fill_year if is_year_round else fill_season
        for cell in row:
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border

    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 44
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"
    ws.sheet_view.showGridLines = False

    table = Table(displayName="CatalogSeasonalityTable", ref=f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def main():
    products = load_products()
    rows = build_rows(products)
    product_ids = {row[1] for row in rows}
    if len(product_ids) != 100:
        raise RuntimeError(f"Expected 100 distinct products, got {len(product_ids)}")

    wb = Workbook()
    ws = wb.active
    ws.title = "Mùa vụ sản phẩm"
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    format_workbook(ws)
    OUTPUT_WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_WORKBOOK)
    print(f"Saved {OUTPUT_WORKBOOK}")
    print(f"Products: {len(product_ids)}")
    print(f"Season rows: {len(rows)}")


if __name__ == "__main__":
    main()

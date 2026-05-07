from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parents[2]
SOURCE_WORKBOOK = REPO_ROOT / "output" / "spreadsheet" / "recommendation_product_data_review_multibatch.xlsx"
OUTPUT_WORKBOOK = REPO_ROOT / "output" / "spreadsheet" / "catalog_product_list.xlsx"
EXTRACTED_AT = "2026-04-24 14:43:30 ICT"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
THIN_BORDER = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)

FIELDS = [
    "ProductID",
    "CategoryName",
    "ProductName",
    "Sku",
    "Current Origin",
    "Current Standard",
    "Current Preservation",
    "Current Weight",
    "Current NearExpiryDays",
    "Current ShortDescription",
    "Current Origin Classification",
    "Research Priority",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", default=str(SOURCE_WORKBOOK))
    parser.add_argument("--output", default=str(OUTPUT_WORKBOOK))
    parser.add_argument("--timestamp", default=EXTRACTED_AT)
    return parser.parse_args()


def auto_fit(ws, max_width: int = 42) -> None:
    for column_cells in ws.columns:
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_len + 2, 12), max_width)


def main() -> None:
    args = parse_args()
    source_wb = load_workbook(Path(args.source), data_only=False)
    source_ws = source_wb["Products Review"]
    headers = [cell.value for cell in source_ws[1]]
    idx = {header: i for i, header in enumerate(headers)}

    wb = Workbook()
    ws = wb.active
    ws.title = "Catalog Products"
    output_headers = FIELDS + ["Thời gian trích xuất"]
    ws.append(output_headers)

    for row in source_ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        ws.append([row[idx[field]] for field in FIELDS] + [args.timestamp])

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    auto_fit(ws, max_width=56)

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


if __name__ == "__main__":
    main()

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK_PATH = REPO_ROOT / "output" / "spreadsheet" / "recommendation_product_data_review_multibatch.xlsx"
DEFAULT_TIMESTAMP = "2026-04-24 14:43:30 ICT"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
THIN_BORDER = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK_PATH))
    parser.add_argument("--timestamp", default=DEFAULT_TIMESTAMP)
    return parser.parse_args()


def normalize_empty(value):
    return None if value == "" else value


def ensure_header_column(ws, before_header: str, new_header: str) -> int:
    headers = [cell.value for cell in ws[1]]
    if new_header in headers:
        return headers.index(new_header) + 1

    insert_at = headers.index(before_header) + 1 if before_header in headers else len(headers) + 1
    ws.insert_cols(insert_at)
    header_cell = ws.cell(row=1, column=insert_at)
    header_cell.value = new_header
    header_cell.fill = HEADER_FILL
    header_cell.font = Font(color="FFFFFF", bold=True)
    header_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_cell.border = THIN_BORDER
    return insert_at


def fill_timestamp_products_review(ws, timestamp: str) -> None:
    headers = [cell.value for cell in ws[1]]
    idx = {header: i + 1 for i, header in enumerate(headers)}
    time_col = idx["Thời gian thu thập"]
    status_col = idx["Review Status"]
    primary_col = idx["Primary Source URL"]

    for row in range(2, ws.max_row + 1):
        status = ws.cell(row=row, column=status_col).value
        primary = ws.cell(row=row, column=primary_col).value
        if ws.cell(row=row, column=time_col).value:
            continue
        if primary or (status and status not in {"Pending", "Ready for batch 1", "Blocked - thiếu ProductInfo seed"}):
            ws.cell(row=row, column=time_col).value = timestamp


def fill_timestamp_evidence(ws, timestamp: str) -> None:
    headers = [cell.value for cell in ws[1]]
    idx = {header: i + 1 for i, header in enumerate(headers)}
    time_col = idx["Thời gian thu thập"]
    product_col = idx["ProductID"]

    for row in range(2, ws.max_row + 1):
        product_id = ws.cell(row=row, column=product_col).value
        if not product_id:
            continue
        if not ws.cell(row=row, column=time_col).value:
            ws.cell(row=row, column=time_col).value = timestamp


def auto_fit(ws, max_width: int = 44) -> None:
    for column_cells in ws.columns:
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_len + 2, 12), max_width)


def main() -> None:
    args = parse_args()
    workbook_path = Path(args.workbook)
    wb = load_workbook(workbook_path)

    ensure_header_column(wb["Products Review"], "Tôi chưa chắc chắn", "Thời gian thu thập")
    ensure_header_column(wb["Seasonality Evidence"], "Tôi chưa chắc chắn", "Thời gian thu thập")
    ensure_header_column(wb["Research Evidence"], "Tôi chưa chắc chắn", "Thời gian thu thập")

    fill_timestamp_products_review(wb["Products Review"], args.timestamp)
    fill_timestamp_evidence(wb["Seasonality Evidence"], args.timestamp)
    fill_timestamp_evidence(wb["Research Evidence"], args.timestamp)

    for sheet_name in ("Products Review", "Seasonality Evidence", "Research Evidence"):
        auto_fit(wb[sheet_name])

    wb.save(workbook_path)


if __name__ == "__main__":
    main()
